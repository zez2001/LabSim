"""Build the evidence report with python-docx; render separately before delivery."""
from pathlib import Path
import re,os
from urllib.parse import urlsplit,unquote,quote
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'docs/lab_cuvette_loading/V0_demo04'
OUTPUT=BASE/'LabSim_V0_demo04_原座装载核查报告_20261020.docx'
FONT='Noto Sans CJK SC'

def font(style,size,bold=None):
    style.font.name=FONT;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
    if bold is not None:style.font.bold=bold

def href(raw):
    u=urlsplit(raw)
    if u.scheme:return raw
    p=(BASE/unquote(u.path)).resolve()
    if not p.exists():raise FileNotFoundError(p)
    if p.is_relative_to(BASE):return Path(os.path.relpath(p,BASE)).as_posix()
    if p.is_relative_to(ROOT):return 'https://github.com/zez2001/LabSim/blob/codex/lab-cuvette-demo04/'+quote(p.relative_to(ROOT).as_posix())
    raise ValueError(raw)

def link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(href(url),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'185A82');pr.append(color);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,text):
    pattern=re.compile(r'(\[[^\]]+\]\([^\n)]+\)|\*\*[^*]+\*\*|`[^`]+`)')
    for t in pattern.split(text):
        if t.startswith('[') and '](' in t:
            label,url=t[1:-1].split('](',1);link(p,label,url)
        elif t.startswith('**'):
            p.add_run(t[2:-2]).bold=True
        elif t.startswith('`'):
            p.add_run(t[1:-1])
        else:p.add_run(t)

def build():
    doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Mm(19);sec.bottom_margin=Mm(18);sec.left_margin=Mm(20);sec.right_margin=Mm(20)
    for name,size,bold in [('Normal',10.5,False),('Title',21,True),('Heading 1',14,True),('Heading 2',12,True),('Caption',9,False),('List Bullet',10.5,False)]:font(doc.styles[name],size,bold)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    n=doc.styles['Normal'].paragraph_format;n.space_after=Pt(5);n.line_spacing=1.15
    for name in ['Heading 1','Heading 2']:
        f=doc.styles[name].paragraph_format;f.space_before=Pt(12);f.space_after=Pt(5);f.keep_with_next=True
    font(doc.styles['Title'],21,True);doc.styles['Title'].paragraph_format.space_after=Pt(10)
    doc.core_properties.title='LabSim V0 demo04 原座装载核查报告';doc.core_properties.author='LabSim project';doc.core_properties.subject='用户更正尺寸后的单仪器场景与实际运行证据'
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=foot.add_run();r.font.size=Pt(9);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
    lines=(BASE/'report.md').read_text().splitlines();i=0;in_code=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('```'):in_code=not in_code;continue
        if in_code:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
            r=p.add_run(line);r.font.name='DejaVu Sans Mono';r.font.size=Pt(8);continue
        if line.startswith('# '):
            doc.add_paragraph('LabSim V0 demo04 原座装载核查报告','Title');continue
        if line.startswith('##'):
            text=re.sub(r'^#+\s*','',line);text=re.sub(r'^\d+\W*','',text);text=re.sub(r'[:：，、。/／_()（）—–-]',' ',text)
            doc.add_paragraph(text,'Heading 1' if not line.startswith('###') else 'Heading 2');continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            parsed=[[x.strip() for x in row.strip('|').split('|')] for row in rows if not re.match(r'^\|?[\s:|-]+\|?$',row)]
            count=len(parsed[0]);tab=doc.add_table(rows=0,cols=count);tab.autofit=False
            widths=([.24,.33,.43] if count==3 else [.5,.5] if count==2 else [1/count]*count)
            for ri,row in enumerate(parsed):
                cells=tab.add_row().cells
                for ci,value in enumerate(row):
                    cells[ci].width=Mm(170*widths[ci]);p=cells[ci].paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
                    inline(p,value)
                    for r in p.runs:r.font.size=Pt(9);r.bold=(ri==0) or r.bold
                    if ri==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EFEFEF');cells[ci]._tc.get_or_add_tcPr().append(sh)
                pr=tab.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
                if ri==0:pr.append(OxmlElement('w:tblHeader'))
            borders=OxmlElement('w:tblBorders')
            for side in ['top','bottom','left','right','insideH','insideV']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'CCCCCC');borders.append(e)
            tab._tbl.tblPr.append(borders);continue
        m=re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$',line)
        if m:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            from PIL import Image
            path=(BASE/unquote(m[2])).resolve();w,h=Image.open(path).size
            width=min(170,100*w/h);r=p.add_run();shape=r.add_picture(str(path),width=Mm(width));shape._inline.docPr.set('descr',m[1]);continue
        style='Caption' if re.match(r'^图\s*\d',line) else 'List Bullet' if line.startswith('- ') else None
        text=line[2:] if line.startswith(('- ','> ')) else line
        p=doc.add_paragraph(style=style);inline(p,text)
    doc.save(OUTPUT);print(OUTPUT)
if __name__=='__main__':build()

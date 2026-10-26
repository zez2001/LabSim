#!/usr/bin/env python3
"""Prepare the independent Project_v1_1 Chinese technical handoff report.

All result text and numbers come from explicit inputs and hashed evidence.
Default operation is read-only validation. Actual authoring requires --author
and --marker-already-run after the caller has run the documents skill marker
exactly once. This program never runs that marker or a renderer itself.
Use only the Codex bundled Python for this script and document authoring.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
from urllib.parse import quote, urlsplit


ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
SCHEMA = 'labsim.Project_v1_1.report_inputs.v1'
STEM = 'LabSim_Project_v1_1_项目发布与验证报告'
TITLE = 'LabSim Project v1 1 项目发布与验证报告'
TOPICS = {'summary', 'executor', 'camera', 'grasp', 'gpt', 'teacher_data', 'release'}
SCOPES = {'diagnostic', 'camera', 'grasp', 'teacher', 'gpt_debug', 'gpt_formal', 'release', 'contract', 'history'}
KINDS = {'diagnostic', 'teacher', 'gpt_debug', 'gpt_formal'}
STATUSES = {'closed', 'not_run', 'blocked_budget', 'incomplete'}
BLOCK_TYPES = {'paragraph', 'heading', 'table', 'figure', 'links', 'code', 'markdown'}


def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def digest(path):
    value = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b''):
            value.update(chunk)
    return value.hexdigest()


def plain_heading(value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError('Title and headings must contain plain descriptive text')
    if any(not (char.isalnum() or char.isspace()) for char in value):
        raise ValueError('Title and headings permit words, numbers and spaces only: ' + value)
    return value


def local_file(value):
    relative = Path(value)
    if relative.is_absolute() or '..' in relative.parts:
        raise ValueError('Evidence must use repository-relative paths without traversal: ' + str(value))
    current = ROOT
    for part in relative.parts:
        current /= part
        if current.is_symlink():
            raise ValueError('Evidence must not traverse a symlink: ' + str(value))
    path = (ROOT / relative).resolve()
    if not path.is_relative_to(ROOT) or not path.is_file():
        raise FileNotFoundError(value)
    with path.open('rb') as stream:
        if stream.read(100).startswith(b'version https://git-lfs.github.com/spec/v1'):
            raise ValueError('Evidence is an unmaterialized LFS pointer: ' + str(value))
    return path


def public_url(value):
    parsed = urlsplit(value)
    if parsed.scheme != 'https' or not parsed.netloc:
        raise ValueError('Only explicit public HTTPS URLs are allowed: ' + str(value))
    return value


def json_pointer(value, pointer):
    if pointer == '':
        return value
    if not isinstance(pointer, str) or not pointer.startswith('/'):
        raise ValueError('Evidence JSON pointer must be empty or start with /')
    for part in pointer[1:].split('/'):
        part = part.replace('~1', '/').replace('~0', '~')
        value = value[int(part)] if isinstance(value, list) else value[part]
    return value


def source_registry(data, *, require_hashes):
    records = {}
    for entry in data.get('sources', []):
        item = dict(entry)
        key = item['id']
        if not re.fullmatch(r'[A-Za-z0-9_-]+', key) or key in records:
            raise ValueError('Source IDs must be distinct simple identifiers: ' + key)
        if item.get('scope') not in SCOPES:
            raise ValueError('Source needs an explicit current evidence scope: ' + key)
        if bool(item.get('path')) == bool(item.get('url')):
            raise ValueError('Source needs exactly one path or HTTPS url: ' + key)
        if item.get('url'):
            item['_target'] = public_url(item['url'])
            if item['scope'] == 'history' and 'github.com/' in item['url'] and not re.search(r'/(blob|tree)/[0-9a-f]{40}/', item['url']):
                raise ValueError('Historical GitHub links must use a fixed complete commit: ' + key)
        else:
            path = local_file(item['path'])
            if path.suffix.lower() not in {'.json', '.md', '.csv', '.yaml', '.yml', '.txt', '.png', '.jpg', '.jpeg', '.mp4'}:
                raise ValueError('Use compact public report evidence instead of raw datasets/archives: ' + key)
            # Historical material is linked with fixed public commits rather
            # than copied into this current-only report's dependencies.
            if item['scope'] == 'history' or re.search(r'/(Project_v1|Size_v1|Env_v1)/', '/' + item['path']):
                raise ValueError('Use a fixed public URL for historical sources: ' + key)
            actual = digest(path)
            expected = item.get('sha256')
            if require_hashes and not re.fullmatch(r'[0-9a-f]{64}', str(expected or '')):
                raise ValueError('Final local evidence needs a saved SHA256: ' + key)
            if expected and expected != actual:
                raise ValueError('Evidence changed since input preparation: ' + key)
            item.update(_path=path, _sha256=actual, _bytes=path.stat().st_size)
            if path.suffix == '.json':
                item['_json'] = load(path)
        records[key] = item
    return records


def text_value(value, records):
    if isinstance(value, dict):
        source = records[value['source']]
        if source['scope'] == 'history':
            raise ValueError('Historical evidence cannot supply current result values')
        if '_json' not in source:
            raise ValueError('Metric references require a local JSON source')
        result = json_pointer(source['_json'], value.get('pointer', ''))
        if result is None:
            # Unknown task outcomes must never silently become false or zero.
            return value.get('null_text', 'null 未判定')
        if isinstance(result, (list, dict)):
            raise ValueError('Metric references must resolve to one scalar')
        if isinstance(result, float) and not math.isfinite(result):
            raise ValueError('Nonfinite metrics require an explicit explanation, not a numeric result')
        if value.get('format'):
            if isinstance(result, bool) or not isinstance(result, (int, float)):
                raise ValueError('Numeric formatting requires a numeric scalar')
            result = format(result, value['format'])
        else:
            result = scalar_text(result)
        return value.get('prefix', '') + str(result) + value.get('suffix', '')
    return scalar_text(value)


def scalar_text(value):
    if value is None:
        return 'null 未判定'
    if value is True:
        return 'true'
    if value is False:
        return 'false'
    if not isinstance(value, (str, int, float)):
        raise ValueError('Body/table values must be strings or scalar evidence references')
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Nonfinite scalars are not report results')
    return str(value)


def markdown_fragment(text):
    """Import an editorial fragment, not a full externally linked report.

    Supports paragraphs, plain headings, bullets and fenced code. Explicit
    table/figure/link blocks keep widths, source hashes and dependencies clear.
    """
    blocks, paragraph, code = [], [], None

    def flush():
        if paragraph:
            blocks.append(dict(type='paragraph', text=' '.join(paragraph)))
            paragraph.clear()

    for line in text.splitlines():
        if line.startswith('```'):
            flush()
            if code is None:
                code = []
            else:
                blocks.append(dict(type='code', text='\n'.join(code)))
                code = None
        elif code is not None:
            code.append(line)
        elif not line.strip():
            flush()
        elif re.match(r'^#{1,3} ', line):
            flush()
            blocks.append(dict(type='heading', text=plain_heading(line.lstrip('#').strip()), level=2))
        elif line.startswith(('- ', '* ')):
            flush()
            blocks.append(dict(type='paragraph', text='• ' + line[2:].strip()))
        elif line.lstrip().startswith('|') or re.search(r'!?\[[^\]]*\]\(', line) or '<a ' in line:
            raise ValueError('Use explicit table/figure/links blocks instead of linked Markdown fragments')
        else:
            paragraph.append(line.strip())
    if code is not None:
        raise ValueError('Unclosed fenced code in Markdown fragment')
    flush()
    return blocks


def resolve_block(block, records):
    kind = block.get('type')
    if kind not in BLOCK_TYPES:
        raise ValueError('Unsupported report block: ' + str(kind))
    item = dict(block)
    if kind in {'paragraph', 'heading', 'code'}:
        item['text'] = text_value(item['text'], records)
        if kind == 'heading':
            plain_heading(item['text'])
        if kind == 'code' and any(len(line) > 90 for line in item['text'].splitlines()):
            raise ValueError('Wrap long command lines before authoring; maximum 90 characters')
    elif kind == 'table':
        item['headers'] = [text_value(value, records) for value in item['headers']]
        item['rows'] = [[text_value(value, records) for value in row] for row in item['rows']]
        count = len(item['headers'])
        if not 2 <= count <= 7 or any(len(row) != count for row in item['rows']):
            raise ValueError('Use 2 to 7 columns and a consistent table schema')
        if not item['rows'] or len(item['rows']) > 18:
            raise ValueError('Keep main report tables concise; link full episode records')
        item['widths'] = [float(value) for value in item.get('widths', [1] * count)]
        if len(item['widths']) != count or any(float(width) <= 0 for width in item['widths']):
            raise ValueError('Positive intentional table widths are required')
        item['align'] = item.get('align', ['left'] * count)
        if len(item['align']) != count or set(item['align']) - {'left', 'center', 'right'}:
            raise ValueError('Table alignment must be declared per column')
        if any(len(cell) > 240 for row in item['rows'] for cell in row):
            raise ValueError('Move narrative explanations out of table cells')
    elif kind == 'figure':
        source = records[item['source']]
        if '_path' not in source or source['_path'].suffix.lower() not in {'.png', '.jpg', '.jpeg'}:
            raise ValueError('Figures need local PNG/JPEG evidence sources')
        item['_path'] = source['_path']
        item['caption'] = text_value(item['caption'], records)
        item['max_height_cm'] = float(item.get('max_height_cm', 6.5))
        if not 2 <= item['max_height_cm'] <= 12:
            raise ValueError('Figure height must preserve readable page layout')
    elif kind == 'links':
        if not item.get('items'):
            raise ValueError('Link groups must not be empty')
        for link in item['items']:
            if link['source'] not in records or not link.get('label'):
                raise ValueError('Links require a known source and descriptive label')
    elif kind == 'markdown':
        source = records[item['source']]
        if '_path' not in source or source['_path'].suffix != '.md' or source['scope'] == 'history':
            raise ValueError('Editorial fragments must be current local Markdown')
        return [resolve_block(child, records)[0] for child in markdown_fragment(source['_path'].read_text(encoding='utf-8'))]
    return [item]


def prepare(data, *, allow_draft=False):
    if data.get('schema') != SCHEMA or data.get('project_version') != 'Project_v1_1':
        raise ValueError('Use an independent Project_v1_1 input schema')
    finalized = data.get('finalized') is True
    if not finalized and not allow_draft:
        raise ValueError('Inputs are not finalized; --allow-draft is validation-only')
    meta = data['metadata']
    plain_heading(meta.get('title', TITLE))
    if not meta.get('conclusion') or not meta.get('date') or not meta.get('audience'):
        raise ValueError('Opening needs the main conclusion, date and intended reader')
    if finalized:
        date.fromisoformat(meta['date'])
    if not data.get('limitations') or any(not isinstance(value, str) or not value for value in data['limitations']):
        raise ValueError('Explicit evidence boundaries and unresolved work are required')
    records = source_registry(data, require_hashes=finalized)
    cohorts = data.get('cohorts', [])
    if {cohort.get('kind') for cohort in cohorts} != KINDS:
        raise ValueError('Declare diagnostic, teacher, GPT debug and GPT formal separately')
    for cohort in cohorts:
        status = cohort.get('status')
        if status not in STATUSES or not cohort.get('note'):
            raise ValueError('Cohorts need explicit execution status and interpretation')
        completed = cohort.get('completed_count', 0)
        if not isinstance(completed, int) or isinstance(completed, bool) or completed < 0:
            raise ValueError('Completed counts must be actual nonnegative integers')
        if status == 'closed':
            if not completed or cohort.get('summary_source') not in records:
                raise ValueError('Closed cohorts require nonzero actual count and summary evidence')
            source_scope = records[cohort['summary_source']]['scope']
            if source_scope != cohort['kind'] and not (cohort['kind'] == 'diagnostic' and source_scope in {'camera', 'grasp'}):
                raise ValueError('A closed cohort must cite its own version and evidence scope')
            if cohort.get('planned_count') is not None and completed != cohort['planned_count']:
                raise ValueError('A closed planned cohort must retain every attempted episode')
        elif status == 'not_run' and completed:
            raise ValueError('not_run must have zero actual completed episodes')
    pages = []
    topics = set()
    for source_page in data.get('pages', []):
        page = dict(source_page)
        plain_heading(page['title'])
        declared = set(page.get('topics', []))
        if not declared or declared - TOPICS:
            raise ValueError('Page groups need supported technical topics')
        topics.update(declared)
        page['blocks'] = [resolved for block in page.get('blocks', []) for resolved in resolve_block(block, records)]
        if not page['blocks']:
            raise ValueError('Do not create empty logical page groups')
        pages.append(page)
    if not 6 <= len(pages) <= 8 or topics != TOPICS or 'summary' not in pages[0]['topics']:
        raise ValueError('Plan 6 to 8 page groups covering summary, executor, camera, grasp, GPT, teacher/data and release')
    return dict(metadata=meta, pages=pages, sources=records, cohorts=cohorts,
                limitations=data['limitations'], finalized=finalized)


def target(source, base):
    if source.get('url'):
        return source['url']
    return quote(os.path.relpath(source['_path'], base), safe='/._-')


def write_markdown(report, path):
    meta, records = report['metadata'], report['sources']
    lines = ['# ' + meta.get('title', TITLE), '',
             f"日期 {meta['date']}　面向 {meta['audience']}", '', meta['conclusion'], '']
    for page in report['pages']:
        lines += ['## ' + page['title'], '']
        for block in page['blocks']:
            kind = block['type']
            if kind == 'heading':
                lines += ['### ' + block['text'], '']
            elif kind == 'paragraph':
                lines += [block['text'], '']
            elif kind == 'code':
                lines += ['```' + block.get('language', 'bash'), block['text'], '```', '']
            elif kind == 'table':
                clean = lambda value: value.replace('|', '/').replace('\n', '<br>')
                lines += ['| ' + ' | '.join(map(clean, block['headers'])) + ' |',
                          '| ' + ' | '.join('---' for _ in block['headers']) + ' |']
                lines += ['| ' + ' | '.join(map(clean, row)) + ' |' for row in block['rows']]
                lines += ['']
            elif kind == 'figure':
                lines += [f"![{block['caption']}]({target(records[block['source']], path.parent)})", '']
            elif kind == 'links':
                lines += [' · '.join(f"[{link['label']}]({target(records[link['source']], path.parent)})" for link in block['items']), '']
    lines += ['### 验证范围与限制', '']
    for value in report['limitations']:
        lines += [value, '']
    path.write_text('\n'.join(lines), encoding='utf-8')


def write_docx(report, path):
    # Lazy imports allow schema/evidence checks without creating an artifact.
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Cm, Inches, Pt, RGBColor
    from PIL import Image

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    body_cm = section.page_width.cm - 4
    for style in doc.styles:
        for border in list(style.element.findall('.//' + qn('w:pBdr'))):
            border.getparent().remove(border)
    for name, size, spacing in [('Normal', 11, 15), ('Title', 20, 26),
                                ('Heading 1', 14, 19), ('Heading 2', 11.8, 16), ('Caption', 9.5, 12.8)]:
        style = doc.styles[name]
        style.font.name = 'DejaVu Sans'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key, value in [('eastAsia', 'Noto Sans CJK SC'), ('ascii', 'DejaVu Sans'), ('hAnsi', 'DejaVu Sans')]:
            fonts.set(qn('w:' + key), value)
        style.paragraph_format.line_spacing = Pt(spacing)
        style.paragraph_format.space_after = Pt(6)
        snap = OxmlElement('w:snapToGrid')
        snap.set(qn('w:val'), '0')
        style.element.get_or_add_pPr().append(snap)
    for name in ('Heading 1', 'Heading 2'):
        doc.styles[name].paragraph_format.space_before = Pt(7)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles['Caption'].font.italic = False
    doc.styles['Caption'].font.bold = False
    meta, records = report['metadata'], report['sources']
    doc.add_paragraph(meta.get('title', TITLE), 'Title')
    doc.add_paragraph(f"日期 {meta['date']}　面向 {meta['audience']}", 'Caption')
    doc.add_paragraph(meta['conclusion'])

    def links(block):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        for index, link in enumerate(block['items']):
            if index:
                paragraph.add_run('　·　')
            hyp = OxmlElement('w:hyperlink')
            hyp.set(qn('r:id'), paragraph.part.relate_to(
                target(records[link['source']], path.parent), RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run, props = OxmlElement('w:r'), OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '24465B'); props.append(color)
            size = OxmlElement('w:sz'); size.set(qn('w:val'), '20'); props.append(size)
            fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC'); props.append(fonts)
            run.append(props)
            text = OxmlElement('w:t'); text.text = link['label']; run.append(text)
            hyp.append(run); paragraph._p.append(hyp)

    def table(block):
        obj = doc.add_table(rows=1, cols=len(block['headers']))
        obj.alignment, obj.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        widths = [body_cm * float(width) / sum(block['widths']) for width in block['widths']]
        for column, grid, width in zip(obj.columns, obj._tbl.tblGrid.gridCol_lst, widths):
            column.width = Cm(width); grid.w = Cm(width)
        table_width = obj._tbl.tblPr.find(qn('w:tblW'))
        table_width.set(qn('w:w'), str(Cm(body_cm).twips)); table_width.set(qn('w:type'), 'dxa')
        for index, value in enumerate(block['headers']):
            obj.rows[0].cells[index].text = value
        for values in block['rows']:
            cells = obj.add_row().cells
            for index, value in enumerate(values):
                cells[index].text = value
        obj.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row_index, row in enumerate(obj.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for index, cell in enumerate(row.cells):
                cell.width = Cm(widths[index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                prop = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ('top', 'left', 'bottom', 'right'):
                    border = OxmlElement('w:' + edge)
                    for key, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
                        border.set(qn('w:' + key), value)
                    borders.append(border)
                prop.append(borders)
                margins = OxmlElement('w:tcMar')
                for edge in ('top', 'left', 'bottom', 'right'):
                    margin = OxmlElement('w:' + edge)
                    margin.set(qn('w:w'), '90'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
                prop.append(margins)
                if row_index == 0 or row_index % 2 == 0:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'E8EEF3' if row_index == 0 else 'F7F8F9')
                    prop.append(shading)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = Pt(12.7)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                                           'right': WD_ALIGN_PARAGRAPH.RIGHT}[block['align'][index]]
                    word_wrap = OxmlElement('w:wordWrap'); word_wrap.set(qn('w:val'), '0')
                    paragraph._p.get_or_add_pPr().append(word_wrap)
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5); run.bold = row_index == 0
                        run.font.color.rgb = RGBColor(0, 0, 0)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.line_spacing = Pt(3)

    for page_index, page in enumerate(report['pages']):
        if page_index:
            doc.add_page_break()
        doc.add_paragraph(page['title'], 'Heading 1')
        for block in page['blocks']:
            kind = block['type']
            if kind in {'paragraph', 'heading'}:
                doc.add_paragraph(block['text'], 'Heading 2' if kind == 'heading' else 'Normal')
            elif kind == 'code':
                paragraph = doc.add_paragraph(block['text'])
                paragraph.paragraph_format.line_spacing = Pt(10.5)
                for run in paragraph.runs:
                    run.font.name = 'DejaVu Sans Mono'; run.font.size = Pt(8.5)
            elif kind == 'table':
                table(block)
            elif kind == 'links':
                links(block)
            elif kind == 'figure':
                with Image.open(block['_path']) as pic:
                    width_px, height_px = pic.size
                width_cm = min(body_cm, block['max_height_cm'] * width_px / height_px)
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.line_spacing = 1
                inline = paragraph.add_run().add_picture(str(block['_path']), width=Cm(width_cm))
                inline._inline.docPr.set('descr', block['caption'])
                doc.add_paragraph(block['caption'], 'Caption')
    doc.add_paragraph('验证范围与限制', 'Heading 2')
    for value in report['limitations']:
        doc.add_paragraph(value)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.core_properties.title = meta.get('title', TITLE)
    doc.core_properties.subject = 'Project_v1_1 实际执行证据与技术交接'
    doc.core_properties.author = meta.get('author', 'LabSim 项目')
    doc.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inputs', required=True, type=Path)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--validate-only', action='store_true', help='Read-only default; never creates outputs')
    mode.add_argument('--author', action='store_true', help='Create MD and DOCX only after the one-time artifact marker')
    parser.add_argument('--allow-draft', action='store_true', help='Permit incomplete finalization metadata for read-only validation')
    parser.add_argument('--marker-already-run', action='store_true', help='Caller confirms the skill marker succeeded exactly once')
    parser.add_argument('--output-dir', type=Path, default=HERE)
    parser.add_argument('--overwrite', action='store_true', help='Replace report outputs while iterating render QA')
    args = parser.parse_args(argv)
    if args.author and (args.allow_draft or not args.marker_already_run):
        parser.error('Authoring requires finalized inputs and --marker-already-run; never author drafts')
    report = prepare(load(args.inputs), allow_draft=args.allow_draft)
    if not args.author:
        print(json.dumps(dict(validated=True, finalized=report['finalized'], planned_page_groups=len(report['pages']),
            evidence_sources=len(report['sources']), authoring_executed=False, artifact_marker_executed=False), ensure_ascii=False))
        return 0
    output = args.output_dir.resolve()
    if not output.is_relative_to(PROJECT) or any(path.is_symlink() for path in (args.output_dir, *args.output_dir.parents)):
        parser.error('Author outputs must remain in the independent Project_v1_1 tree without symlinks')
    targets = [output / (STEM + suffix) for suffix in ('.md', '.docx')]
    build_path = output / 'report_build.json'
    if not args.overwrite and any(path.exists() for path in [*targets, build_path]):
        raise FileExistsError('Preserve existing reports or explicitly use --overwrite for QA revisions')
    output.mkdir(parents=True, exist_ok=True)
    write_markdown(report, targets[0])
    write_docx(report, targets[1])
    records = [dict(id=key, scope=item['scope'], **(
        dict(path=item['path'], report_relative_target=target(item, output), sha256=item['_sha256'], bytes=item['_bytes'])
        if '_path' in item else dict(url=item['url']))) for key, item in report['sources'].items()]
    build = dict(schema='labsim.Project_v1_1.report_build.v1', created_at_utc=datetime.now(timezone.utc).isoformat(),
        inputs_sha256=digest(args.inputs), generator_sha256=digest(__file__), planned_page_groups=len(report['pages']),
        cohort_declarations=report['cohorts'], evidence_sources=records,
        outputs=[dict(path=path.name, bytes=path.stat().st_size, sha256=digest(path)) for path in targets],
        render_inspection_required=True, visual_QA_completed=False,
        actual_rendered_page_count=None, expected_rendered_page_range=[6, 8],
        marker_execution='Caller precondition; this generator never invokes or repeats the artifact marker',
        layout='Preserve Project_v1_1/report and each report_relative_target; DOCX links are not rewritten by package_agent_results.py')
    build_path.write_text(json.dumps(build, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(build['outputs'], ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

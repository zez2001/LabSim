#!/usr/bin/env python3
"""Build the concise Project_v1 Chinese release report from closed evidence.

Run with the Codex bundled Python. The documents artifact marker is an explicit
one-time precondition before the first authoring run; this script never repeats
it. --validate-only reads and assembles evidence without creating deliverables.
Render the resulting DOCX with the unchanged bundled document renderer through
Env_v1/report/render_docx_isolated.py, then inspect every rendered page.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
from urllib.parse import quote, urlsplit

ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
TITLE = "LabSim Project v1 项目发布与验证报告"
STEM = "LabSim_Project_v1_项目发布与验证报告"


def load(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def number(value, digits=2):
    if value is None:
        return "未记录"
    return f"{float(value):.{digits}f}"


def relative(path, base=HERE):
    path = Path(path)
    if not path.is_absolute():
        path = ROOT / path
    return quote(os.path.relpath(path, base), safe="/._-")


def local_file(path):
    """Validate every report dependency as a portable repository file."""
    value = Path(path)
    if value.is_absolute() or ".." in value.parts:
        raise ValueError("Report dependencies must be repository-relative paths: " + str(path))
    current = ROOT
    for part in value.parts:
        current = current / part
        if current.is_symlink():
            raise ValueError("Report dependency must not traverse a symlink: " + str(path))
    target = (ROOT / value).resolve()
    if not target.is_relative_to(ROOT) or not target.is_file():
        raise FileNotFoundError(path)
    with target.open("rb") as stream:
        if stream.read(100).startswith(b"version https://git-lfs.github.com/spec/v1"):
            raise ValueError("Fetch the real report dependency before authoring: " + str(path))
    return target


def dependencies(pages):
    records = {}
    for page in pages:
        for kind, value in page:
            if kind == "image":
                path = Path(value["path"])
            elif kind == "link" and value.get("path"):
                path = local_file(value["path"])
            else:
                continue
            key = path.relative_to(ROOT).as_posix()
            if key in records:
                continue
            digest = hashlib.sha256()
            with path.open("rb") as stream:
                for block in iter(lambda: stream.read(1024 * 1024), b""):
                    digest.update(block)
            records[key] = {"repository_path": key, "report_relative_target": relative(path),
                            "bytes": path.stat().st_size, "sha256": digest.hexdigest()}
    return list(records.values())


def paragraph(text):
    return ("p", text)


def table(header, rows, widths):
    return ("table", {"header": header, "rows": rows, "widths": widths})


def assemble(data):
    if data.get("finalized") is not True:
        raise ValueError("Release inputs must declare finalized=true after all reported runs have closed")
    if len(data.get("gpt_formal", [])) != 5:
        raise ValueError("Keep all five formal GPT attempts in gpt_formal")
    if not re.fullmatch(r"[0-9a-f]{40}", data.get("runtime_commit", "")):
        raise ValueError("runtime_commit must be a verified complete Git SHA")
    if not re.fullmatch(r"[0-9a-f]{64}", data.get("formal_fingerprint", "")):
        raise ValueError("formal_fingerprint must be a complete saved SHA256")
    for item in data.get("evidence", []):
        local_file(item["path"])
    for item in data.get("media", []):
        local_file(item["path"])

    assets = load(PROJECT / "assets/anonymous_download_summary.json")
    sample = load(PROJECT / "data/public_sample_publication.json")
    teacher = data["teacher"]
    runtime = data["runtime_commit"]
    rows = [
        ["teacher GUI", "1/1", "105.82", "851.96", "6348", "完整装载并稳定3秒"],
        ["teacher headless", teacher["result"], teacher["sim_seconds"], teacher["wall_seconds"], teacher["control_steps"], teacher["note"]],
        ["新受控失败", data["controlled_failure"]["result"], data["controlled_failure"]["sim_seconds"], data["controlled_failure"]["wall_seconds"], data["controlled_failure"]["control_steps"], data["controlled_failure"]["note"]],
        ["GPT 正式5次", data["gpt_summary"]["result"], data["gpt_summary"]["sim_seconds"], data["gpt_summary"]["wall_seconds"], data["gpt_summary"]["control_steps"], data["gpt_summary"]["note"]],
        ["π0.5 / Harness", "各0/1", "各0.25", "6.48 / 86.92", "各15", "各5个模型动作，尚未抓取"],
    ]
    page1 = [
        ("title", TITLE),
        paragraph("发布日期 " + data["date"] + "　分支 codex/labsim　面向组内使用与后续策略评测"),
        paragraph(data["conclusion"]),
        ("h", "当前正式任务"),
        paragraph("默认入口为 lab_cuvette_loading / Project_v1 / standard。Franka 从完整白色6×6来源架的A1取出10×10×67.5 mm尺寸假设器皿，放入单台仪器内的原始A座，完成插入、支撑交接、开爪、退让与至少3秒独立稳定。质量20 g；布局、初始化、资产和场景随机化均冻结。原座及历史V0、Env_v1、Size_v1保持独立。"),
        ("h", "本轮实际执行结果"),
        table(["后端与范围", "成功数", "仿真秒", "墙钟秒", "控制步", "结论"], rows, [3.0, 1.4, 1.7, 2.1, 1.5, 7.3]),
        paragraph("teacher使用真值控制，GPT读取两路图像、机器人自身状态、相机标定和允许的指端接触反馈。各后端共用任务、记录器与独立评分；结束请求、正常退出或接口连通不代表成功。墙钟采用episode结果值，π0.5/Harness为适配回路；含启动的进程时间另见执行证明。"),
    ]
    if data.get("poster"):
        image_path = local_file(data["poster"])
        page1 += [("image", {"path": str(image_path), "caption": data["poster_caption"], "max_height_cm": 5.1})]
    for media in data.get("media", [])[:3]:
        page1 += [("link", {"label": media["label"], "path": media["path"]})]

    setup = """GIT_LFS_SKIP_SMUDGE=1 git clone --branch codex/labsim \\
  https://github.com/zez2001/LabSim.git
cd LabSim
export LABSIM_PYTHON=/absolute/path/to/isaac-python
export LABSIM_GEOMETRY_SITE=/absolute/path/to/geometry/site-packages
./labsim assets download
./labsim doctor
./labsim demo --task lab_cuvette_loading --scene standard \\
  --backend teacher --gui --output outputs/project_demo_001
./labsim eval --task lab_cuvette_loading --scene standard \\
  --backend teacher --episodes 5 --headless \\
  --output outputs/project_teacher5_001"""
    setup = setup.replace("cd LabSim\n", "cd LabSim\ngit checkout " + runtime + "\n")
    extra = """./labsim collect --task lab_cuvette_loading --scene standard \\
  --backend teacher --headless --output outputs/project_collect_001
./labsim dataset-check --path outputs/project_collect_001/episode_0000 \\
  --sample-filter success --policy both --output outputs/read_001.json
./labsim collect --task lab_cuvette_loading --scene standard \\
  --backend teacher --controlled-stop-depth-mm 5 --headless \\
  --output outputs/project_failure_001
for attempt in 0000 0001 0002 0003 0004; do
  ./labsim eval --task lab_cuvette_loading --scene standard \\
    --backend gpt6 --episodes 1 --headless --max-calls 160 \\
    --budget-ledger outputs/cubence_budget.json \\
    --output "outputs/gpt5_${attempt}"
done"""
    page2 = [
        ("h", "从干净 checkout 开始使用"),
        paragraph("已验证环境为Linux、Isaac Sim 5.1.0.0、Python 3.11.16、Torch 2.7.0+cu126、NumPy 1.26.0、h5py 3.15.1、OmegaConf 2.3.0与RTX 4090。复用本机已安装的Isaac环境；Shapely 2.2.0由独立site-packages注入，未修改Isaac或既有模型环境。原生task、factory、controller和collector已注册，./labsim与正常任务入口共享实现；安装细节见仓库接入指南。"),
        ("code", setup),
        paragraph("headless演示将 --gui 换为 --headless。输出目录必须全新；下载器默认使用用户缓存，也可设置LABSIM_ASSET_CACHE。doctor检查真实依赖、路径和资产状态，但CPU预检通过不能代替GPU物理验收。"),
        ("h", "采集 读取与 GPT 评测"),
        ("code", extra),
        paragraph("先在本地终端设置CUBENCE_API_KEY。provider固定为Cubence、https://api.cubence.ai/v1、Responses API与gpt-6-astra。探针、调试、重试和正式评测共享预算ledger；已有ledger不应删除以重置预算。错误类别包括任务失败、模型服务失败、协议错误、预算结束与仿真异常。"),
        ("h", "同机新目录复现范围"),
        paragraph(data["clean_checkout"]),
        paragraph(data["validation_summary"]),
    ]

    attempt_rows = []
    for group, label in ((data.get("gpt_pilots", []), "调试"), (data["gpt_formal"], "正式")):
        for index, run in enumerate(group, 1):
            attempt_rows.append([f"{label}{index:02d}", str(run["requests"]), str(run["actions"]), number(run["sim_seconds"]), number(run["wall_seconds"], 1), run["outcome"]])
    page3 = [
        ("h", "GPT 直接控制与全部尝试"),
        paragraph(data["gpt_capabilities"]),
        paragraph(data.get("gpt_freeze", "正式运行使用保存的固定模型、提示词、工具、动作限制与评分版本；调试尝试另列。")),
        paragraph("模型只用通用末端、夹爪、等待、观察和finish工具，不使用专家宏或teacher轨迹。动作绑定episode、observation和action ID；越界/过期指令拒绝，重复ID查询回执。请求期间暂停物理，完整解析验证后才执行。"),
        table(["尝试", "请求", "动作", "仿真秒", "墙钟秒", "结果与最远阶段"], attempt_rows, [1.7, 1.2, 1.2, 1.6, 1.9, 9.4]),
        paragraph("请求数含失败与重试；动作列为实际执行指令，含部分执行，观察/结束与拒绝另计。墙钟为episode结果值，模型等待时物理暂停。调试尝试未并入正式5次结果。"),
        paragraph(data["gpt_classification"]),
        paragraph(data["gpt_budget"]),
        ("h", "动作约定与公平性"),
        paragraph("GPT末端动作以米、弧度、wxyz四元数记录，world/TCP与绝对/增量语义显式，夹爪为双指总开口。单次上限0.04 m/0.15 rad、0.05至3秒，速度上限0.02 m/s与0.15 rad/s。每指上限8 N；运动默认8 N，夹爪工具可显式选0至8 N，不按teacher阶段或真值自动调力。记录值是驱动上限，不是实测接触力。ACT/DP为7臂关节绝对目标加单指位置，9维实际双指命令另存。"),
    ]

    page4 = [
        ("h", "公开资产与许可边界"),
        paragraph("公开HF Dataset为Txie000/SLAI-Labsim。运行资产包含器皿原模型、35个横向尺寸派生体、仪器、来源架、配置元数据、来源及参数假设；视觉和碰撞对应、单位、坐标轴、根节点、质量、质心、惯量、摩擦与逐文件哈希见manifest和资产锁。公开USD仅清除非运行文档字段中的作者目录，43个USD语义比较通过，历史源文件与物理参数未被改写。"),
        table(["发布内容", "固定 revision", "文件与大小", "实际验证"], [
            ["运行资产", assets["revision"][:12], f"HF任务文件82个；加外部机器人共{assets['files_verified']}个 / {assets['bytes_verified']:,} B", "独立无凭据进程、空缓存83/83；离线复用83/83"],
            ["完整成功缩图样例", sample["sample_revision"][:12], f"5个文件 / {sample['sample_total_bytes']:,} B", "匿名下载5/5；正常ACT和DP reader通过"],
            ["样例发布后仓库总量", sample["sample_revision"][:12], f"{sample['repository_file_count']}个 / {sample['repository_total_bytes']:,} B", "增量发布，保留旧发布内容"],
        ], [3.1, 5.4, 4.3, 4.2]),
        paragraph("项目所有者已确认所提供CAD与派生体的公开再分发授权，使用LicenseRef声明；该授权不等同于制造商许可。来源架与项目数值元数据按MIT声明。Franka文件许可仍未核实，因此没有重新上传HF；下载器从现有公开项目Git LFS固定提交取回相同哈希文件，来源与限制明确保留。"),
        ("h", "新旧数据与可移植性"),
        paragraph("历史20条成功与20条受控提前停止失败继续使用原Size_v1数据和正常ACT/DP读取器。本轮读取验证均通过，未重采整批40条、未改历史标签或指纹。失败数据是预设深度停止，不能解释为自然卡滞率。历史索引迁移需保留collection相对目录树。新Project_v1另采完整成功与受控失败，分别关闭并重新读取，详细结果见对应JSON。"),
        paragraph("公开样例来自本轮teacher GUI完整成功，保留全部6348个转换与末端观测，双视角256×256 RGB仅空间缩为96×96；没有时间抽样。动作、状态、事件、标签及时间戳数值保持不变，路径元数据规范为逻辑标识。H5无外部或虚拟数据依赖，正常reader按实际图像尺寸生成batch。96像素样例并非原生256像素训练数据，也不含高频物理JSONL和全部媒体旁文件。"),
        paragraph(data["new_data"]),
    ]

    version_rows = [
        ["GPT正式代码", runtime],
        ["teacher物理代码", data.get("teacher_runtime_commit", runtime)],
        ["任务与评分", "Project_v1.standard.v1 / Project_v1.independent.v1"],
        ["动作接口", "Project_v1.actions.v1"],
        ["HF运行资产", assets["revision"]],
        ["HF成功样例", sample["sample_revision"]],
        ["GUI成功代码", "d09c5ffce2018702702eb0d4c77d862f1eae6faa"],
        ["GUI场景指纹", "75d288fe67e88c749cb578f36f3c246664b3e26d8455926bffbccf3576964de6"],
        ["正式场景指纹", data["formal_fingerprint"]],
        ["GPT控制器指纹", data.get("gpt_controller_sha256", "见正式评测证据")],
    ]
    page5 = [
        ("h", "版本与关键证据入口"),
        table(["字段", "版本或完整标识"], version_rows, [3.6, 13.4]),
        paragraph("原座局部原点与装载轴分别记录：装载目标使用原座坐标加[0.00065, 0.000001, 0] m，不再把局部原点[0,0]误写成目标轴。控制对准预算、判分容差与物理参数分开保存。评分要求真实抓取离架、搬运、目标支撑、脱指、至少0.04 m退让与3秒稳定；稳定窗位置漂移阈值0.0002 m、朝向变化1度，工程阈值不代表实物精度。"),
        paragraph("侧视与腕视名称直接进入正式配置和ACT/DP shape_meta。两相机采集640×480、60 Hz，正式训练读取RGB为256×256，GPT传输图像另按provider冻结的预处理配置；位姿、内外参、时戳和原始相机元数据随episode记录。GUI成功H5的6348条动作、图像、状态与后继时间对齐，8维训练动作与实际9维对称双指命令核验通过。"),
    ]
    for item in data["evidence"]:
        page5 += [("link", {"label": item["label"], "path": item["path"]})]
    page5 += [
        ("h", "剩余问题与后续工作"),
        paragraph(data["remaining"]),
        paragraph("本轮为同机新目录复现，尚未验证第二台机器。器皿尺寸、20 g质量、摩擦、来源架尺寸及机器人动力学仍是仿真假设，real_calibrated=false。保留原生速度与位姿差分不一致等告警；原座完整PhysX烹饪表示尚未全部读回。输入几何抽样和实际接触证据不能证明严格全局穿透上界，亦非0.1 mm真实精度认证。"),
        paragraph(data["delivery_status"]),
        ("link", {"label": "GitHub 发布分支", "url": "https://github.com/zez2001/LabSim/tree/codex/labsim"}),
        ("link", {"label": "Hugging Face 公开资产与样例", "url": "https://huggingface.co/datasets/Txie000/SLAI-Labsim"}),
    ]
    return [page1, page2, page3, page4, page5]


def link_target(value):
    if value.get("url"):
        parsed = urlsplit(value["url"])
        if parsed.scheme not in ("https", "http") or not parsed.netloc:
            raise ValueError("Only explicit public HTTP(S) links are supported")
        return value["url"]
    return relative(local_file(value["path"]))


def write_markdown(pages, path):
    lines = []
    for page in pages:
        for kind, value in page:
            if kind in ("title", "h"):
                lines += [("# " if kind == "title" else "## ") + value, ""]
            elif kind == "code":
                lines += ["```bash", value, "```", ""]
            elif kind == "table":
                lines += ["| " + " | ".join(value["header"]) + " |", "| " + " | ".join("---" for _ in value["header"]) + " |"]
                lines += ["| " + " | ".join(str(x).replace("|", "/") for x in row) + " |" for row in value["rows"]]
                lines += [""]
            elif kind == "image":
                lines += [f"![{value['caption']}]({relative(value['path'])})", ""]
            elif kind == "link":
                lines += [f"[{value['label']}]({link_target(value)})", ""]
            else:
                lines += [value, ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_docx(pages, path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from PIL import Image

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.left_margin = section.right_margin = Cm(1.9)
    for style in doc.styles:
        for border in list(style.element.findall(".//" + qn("w:pBdr"))):
            border.getparent().remove(border)
    for name, size in (("Normal", 10.1), ("Title", 20), ("Heading 1", 13.8), ("Caption", 8.8)):
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.paragraph_format.line_spacing = Pt({"Normal": 13.3, "Title": 25, "Heading 1": 17.5, "Caption": 11.4}[name])
        style.paragraph_format.space_after = Pt(5)
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        style.element.get_or_add_pPr().append(snap)
    doc.styles["Caption"].font.bold = False
    doc.styles["Caption"].font.italic = False
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(6)

    for index, page in enumerate(pages):
        if index:
            doc.add_page_break()
        for kind, value in page:
            if kind in ("title", "h"):
                doc.add_paragraph(value, "Title" if kind == "title" else "Heading 1")
            elif kind == "table":
                obj = doc.add_table(rows=1, cols=len(value["header"]))
                obj.alignment = WD_TABLE_ALIGNMENT.CENTER
                obj.autofit = False
                for column, grid, width in zip(obj.columns, obj._tbl.tblGrid.gridCol_lst, value["widths"]):
                    column.width = Cm(width)
                    grid.w = Cm(width)
                tw = obj._tbl.tblPr.find(qn("w:tblW"))
                tw.set(qn("w:w"), str(sum(Cm(w).twips for w in value["widths"])))
                tw.set(qn("w:type"), "dxa")
                for j, text in enumerate(value["header"]):
                    obj.rows[0].cells[j].text = text
                for row in value["rows"]:
                    cells = obj.add_row().cells
                    for j, text in enumerate(row):
                        cells[j].text = str(text)
                obj.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                for i, row in enumerate(obj.rows):
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    for j, cell in enumerate(row.cells):
                        cell.width = Cm(value["widths"][j])
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        prop = cell._tc.get_or_add_tcPr()
                        borders = OxmlElement("w:tcBorders")
                        for edge in ("top", "left", "bottom", "right"):
                            border = OxmlElement("w:" + edge)
                            for key, val in {"val": "single", "sz": "5", "color": "D9D9D9"}.items():
                                border.set(qn("w:" + key), val)
                            borders.append(border)
                        prop.append(borders)
                        margins = OxmlElement("w:tcMar")
                        for edge in ("top", "left", "bottom", "right"):
                            el = OxmlElement("w:" + edge)
                            el.set(qn("w:w"), "75")
                            el.set(qn("w:type"), "dxa")
                            margins.append(el)
                        prop.append(margins)
                        if i == 0:
                            fill = OxmlElement("w:shd")
                            fill.set(qn("w:fill"), "E6EDF3")
                            prop.append(fill)
                        for para in cell.paragraphs:
                            para.paragraph_format.space_after = Pt(0)
                            para.paragraph_format.line_spacing = Pt(11.4)
                            if len(value["header"]) >= 5 and j in (1, 2, 3, 4):
                                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in para.runs:
                                run.font.size = Pt(8.7)
                                run.bold = i == 0
                gap = doc.add_paragraph()
                gap.paragraph_format.line_spacing = Pt(3)
                gap.paragraph_format.space_after = Pt(3)
            elif kind == "code":
                p = doc.add_paragraph(value)
                p.paragraph_format.line_spacing = 1.03
                for run in p.runs:
                    run.font.name = "DejaVu Sans Mono"
                    run.font.size = Pt(7.5)
            elif kind == "image":
                with Image.open(value["path"]) as pic:
                    w, h = pic.size
                width = min(17.0, value["max_height_cm"] * w / h)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(value["path"], width=Cm(width))
                doc.add_paragraph(value["caption"], "Caption")
            elif kind == "link":
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                hyp = OxmlElement("w:hyperlink")
                relation = p.part.relate_to(link_target(value), RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                hyp.set(qn("r:id"), relation)
                run = OxmlElement("w:r")
                props = OxmlElement("w:rPr")
                size = OxmlElement("w:sz")
                size.set(qn("w:val"), "18")
                props.append(size)
                run.append(props)
                text = OxmlElement("w:t")
                text.text = value["label"]
                run.append(text)
                hyp.append(run)
                p._p.append(hyp)
            else:
                doc.add_paragraph(value)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Project_v1 公开资产 同机复现 统一策略接口与真实闭环结果"
    doc.core_properties.author = "LabSim 项目"
    doc.save(path)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--inputs", type=Path, required=True)
    p.add_argument("--validate-only", action="store_true")
    args = p.parse_args()
    data = load(args.inputs)
    pages = assemble(data)
    links = dependencies(pages)
    if args.validate_only:
        print(json.dumps({"assembled_pages": len(pages), "evidence_files": len(data["evidence"]), "portable_link_dependencies": len(links), "authoring_executed": False}, ensure_ascii=False))
        return
    write_markdown(pages, HERE / (STEM + ".md"))
    write_docx(pages, HERE / (STEM + ".docx"))
    outputs = []
    for suffix in (".md", ".docx"):
        path = HERE / (STEM + suffix)
        outputs.append({"path": path.name, "bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    (HERE / "report_build.json").write_text(json.dumps({"schema": "labsim.Project_v1.report_build.v1", "inputs_sha256": hashlib.sha256(args.inputs.read_bytes()).hexdigest(), "planned_pages": len(pages), "outputs": outputs, "link_dependencies": links, "package_layout_requirement": "Preserve report/ and each report_relative_target under one Project_v1 parent. Markdown may be relocated by the package helper; DOCX hyperlink targets are not rewritten.", "render_inspection_required": True}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(outputs, ensure_ascii=False))


if __name__ == "__main__":
    main()

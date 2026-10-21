#!/usr/bin/env python3
"""Build a current Size_v1 report from actual evidence, using bundled Python.

Root must run the documents skill marker immediately before first authoring.
--validate-only performs source checks and content assembly without creating
DOCX/Markdown or running the marker. Render only with the packaged renderer;
every rendered page must be visually inspected before delivery.
"""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
from datetime import datetime
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[4]
GEOMETRY = 'docs/lab_cuvette_loading/Size_v1/geometry/family_20261021_02'
COOKED = 'docs/lab_cuvette_loading/Size_v1/physics/cooked_preview01_summary.json'
FAILURE1050 = 'docs/lab_cuvette_loading/Size_v1/data/loading1050_inspection/failure_audit.json'
BUNDLED_PYTHON = '/home/wujing/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3'


def get(value, dotted, default=None):
    for part in dotted.split('.'):
        if not isinstance(value, dict) or part not in value:
            return default
        value = value[part]
    return value


def state(value):
    return '通过' if value is True else '未通过' if value is False else '未知或未执行'


def num(value, digits=2):
    return '未知' if value is None else f'{float(value):.{digits}f}'


def short(value, length=46):
    text = str(value) if value is not None else '未知'
    return text if len(text) <= length else text[:length-1] + '…'


def optional_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and value in (0, 1):
        return bool(value)
    if isinstance(value, str) and value.lower() in ('true', 'false'):
        return value.lower() == 'true'
    return None


def backend_snapshot(metadata, configuration, builder_source=''):
    """Separate optional effective metadata from unchanged builder declarations."""
    sections = [metadata]
    sections.extend(metadata[key] for key in ('runtime_physics', 'physics_configuration',
        'effective_physics', 'effective_settings', 'backend') if isinstance(metadata.get(key), dict))
    def first(keys):
        return next((section[key] for section in sections for key in keys
                     if key in section and not isinstance(section[key], dict)), None)
    gpu = optional_bool(first(('gpu_dynamics_enabled', 'enable_gpu_dynamics', 'enableGPUDynamics')))
    backend = first(('physics_backend', 'backend_name', 'backend'))
    ccd = optional_bool(first(('ccd', 'enable_ccd', 'enableCCD')))
    result = {'backend': backend, 'gpu_dynamics': gpu, 'ccd': ccd,
        'backend_source': 'runtime_metadata' if backend is not None or gpu is not None else 'unknown',
        'ccd_source': 'runtime_metadata' if ccd is not None else 'unknown',
        'broadphase': first(('broadphase', 'broadphase_type', 'broadphaseType'))}
    if backend is None and gpu is None and 'CreateEnableGPUDynamicsAttr(False)' in builder_source:
        result.update(backend='PhysX', gpu_dynamics=False, backend_source='shared_builder_code_declaration')
        if "CreateBroadphaseTypeAttr('MBP')" in builder_source:
            result['broadphase'] = 'MBP'
    if ccd is None:
        result['ccd'] = optional_bool(get(configuration, 'physics.ccd'))
        result['ccd_source'] = 'saved_configuration' if result['ccd'] is not None else 'unknown'
    return result


def failure1050_paragraphs(audit):
    """Render the audited historical failure without relabelling its future runs."""
    if not audit:
        return ['尚未提供10.50mm开发运行的独立失败复核；具体成因保持未知。']
    worst = audit.get('maximum_position_gate_observation', {})
    support = audit.get('support_stage_additional_gate', {})
    xy = support.get('xy_um_range')
    bounds = (num(xy[0], 3) + '至' + num(xy[1], 3)) if isinstance(xy, list) and len(xy) == 2 else '未知'
    verdict = audit.get('original_independent_verdict', {})
    online = get(audit, 'controller_vs_independent_gate.controller_online_safety_origin_drift_limit_m')
    online_um = None if online is None else online * 1e6
    anchor = get(audit, 'grip_thresholds.configured_grasp_anchor_point_m')
    anchor_distance_mm = sum(float(x) ** 2 for x in anchor) ** .5 * 1000 if isinstance(anchor, list) and len(anchor) == 3 else None
    return [
        f"10.50mm开发运行在GRASP阶段{num(worst.get('t_s'), 3)}秒观察到杯底原点相对手漂移{num(worst.get('origin_drift_um'), 3)}µm，"
        f"相对姿态变化{num(worst.get('relative_rotation_change_deg'), 6)}度；配置抓取锚点漂移{num(worst.get('configured_grasp_anchor_drift_um'), 3)}µm。"
        f"杯底原点距配置抓取锚点{num(anchor_distance_mm, 1)}mm，旋转杠杆分量为{num(worst.get('rotation_lever_component_um'), 3)}µm。不能将原点漂移解释为等量纯轴向滑移，也不能用较小的锚点值替换原{num(get(audit, 'grip_thresholds.origin_drift_um'), 1)}µm判分门槛。"
        f"控制器在线安全上限为{num(online_um, 1)}µm，独立质量门控与在线停止条件分别记录；动作继续不等于评分通过。",
        f"第二个独立门控同样未通过：获得承托时的XY误差为{bounds}µm，高于原{num(get(audit, 'alignment_thresholds.xy_um'), 1)}µm门槛；"
        f"{support.get('samples', '未知')}个承托样本中，满足装载对准的样本为{support.get('aligned_supported_samples', '未知')}。"
        '因此不能只忽略早期抓持门控就把该次改判为合格装载。',
        f"该次实际松爪{state(verdict.get('release_observed'))}、撤离{state(verdict.get('retreat_observed'))}、末尾稳定{state(get(verdict, 'post_retreat_window.stable'))}，"
        f"完整任务仍为{state(verdict.get('task_success'))}，失败类型{verdict.get('failure_type', '未知')}。"
        f"进程实际退出码{get(audit, 'process_status.exit_code', '未知')}单独保留，不把进程状态当作上述物理失败成因。"
        '10.50mm开发资产与最终冻结10.00mm场景不同，未获正式配额资格；释放和末态稳定不自动使它成为合格成功或配额失败。'
    ]


class Evidence:
    def __init__(self, root=ROOT):
        self.root = Path(root).resolve()
        self.sources, self.warnings = {}, []

    def path(self, value, relative_to=None):
        p = Path(value)
        if p.is_absolute():
            return p.resolve()
        root_path = self.root / p
        if root_path.exists() or relative_to is None:
            return root_path.resolve()
        return (Path(relative_to) / p).resolve()

    def register(self, value, *, required=True, relative_to=None):
        p = self.path(value, relative_to)
        if not p.is_file():
            if required:
                raise FileNotFoundError('Report evidence missing: ' + str(p))
            return None
        digest = hashlib.sha256()
        with p.open('rb') as stream:
            for block in iter(lambda: stream.read(1 << 20), b''):
                digest.update(block)
        key = str(p.relative_to(self.root)) if p.is_relative_to(self.root) else str(p)
        self.sources[key] = {'path': key, 'bytes': p.stat().st_size, 'sha256': digest.hexdigest()}
        return p

    def json(self, value, *, required=True):
        p = self.register(value, required=required)
        return json.loads(p.read_text(encoding='utf-8')) if p else {}

def _media_items(e, manifest_path):
    if not manifest_path:
        return [], {}
    manifest_path = e.register(manifest_path)
    body = e.json(manifest_path)
    raw = body if isinstance(body, list) else []
    if isinstance(body, dict):
        for key in ('files', 'media', 'videos', 'figures'):
            if isinstance(body.get(key), list):
                raw += body[key]
    items, seen = [], set()
    for item in raw:
        item = {'path': item} if isinstance(item, str) else dict(item)
        source = item.get('path') or item.get('output_path') or item.get('derived_path')
        if not source:
            continue
        path = e.register(source, relative_to=manifest_path.parent)
        if path.suffix.lower() not in {'.mp4', '.webm', '.png', '.jpg', '.jpeg'}:
            continue
        if str(path) in seen:
            continue
        seen.add(str(path))
        declared = item.get('sha256') or item.get('output_sha256')
        if declared and hashlib.sha256(path.read_bytes()).hexdigest() != declared:
            raise ValueError('Declared selected media SHA differs: ' + str(path))
        item.update(path=str(path), label=item.get('label') or item.get('title') or path.name,
            kind='video' if path.suffix.lower() in {'.mp4', '.webm'} else 'image')
        items.append(item)
    return items, body


def collect(e, args):
    summary_path = e.register(args.summary)
    summary = e.json(summary_path)
    if summary.get('schema') != 'labsim.Size_v1.combined_results.v1':
        raise ValueError('Use the actual summarize_lab_cuvette_size_v1.py output')
    if any('complete_task_result_scope' not in row for row in summary.get('candidates', [])):
        raise ValueError('Regenerate the summary with natural-control-only size verdicts before report authoring')
    geometry_path = e.path(args.geometry or GEOMETRY)
    geometry = e.json(geometry_path / 'family_summary.json')
    registry = e.json('assets/lab_cuvette_loading/Size_v1/registry.json')
    cooked = e.json(COOKED, required=False)
    failure1050 = e.json(FAILURE1050, required=False)
    quota = summary.get('collection20_success20_failure', {})
    audit_path = args.collection_audit or get(quota, 'qualification_evidence.path')
    audit = e.json(audit_path) if audit_path else {}
    if audit and audit.get('schema') != 'labsim.Size_v1.collection_audit.v1':
        raise ValueError('Collection evidence must be the existing validator result')
    if audit and get(quota, 'qualification_evidence.sha256'):
        if hashlib.sha256(e.path(audit_path).read_bytes()).hexdigest() != get(quota, 'qualification_evidence.sha256'):
            raise ValueError('Summarized qualification audit differs; regenerate summary first')
    if audit and not get(quota, 'qualification_evidence.sha256'):
        raise ValueError('Regenerate the actual summary with this collection audit before report authoring')
    runs = summary.get('runs', [])
    closed = [r for r in runs if r.get('complete') is True and r.get('task_type') == 'loading'
              and 'preview_or_interface_diagnostic' not in r.get('qualification_exclusions', [])]
    if args.primary_episode_id:
        primary = next((r for r in runs if r.get('episode_id') == args.primary_episode_id), None)
        if primary is None:
            raise ValueError('Requested primary episode ID is absent from the actual summary')
    else:
        primary = next((r for r in closed if r.get('task_success') is True), closed[-1] if closed else None)
    result, config = {}, {}
    if primary:
        result = e.json(primary['result_path']) if primary.get('result_path') else {}
        config = primary.get('config_snapshot') or {}
    if not config:
        config = next((row['config_snapshot'] for row in reversed(runs)
            if isinstance(row.get('config_snapshot'), dict) and row['config_snapshot']), {})
    if args.actual_config_json:
        config = e.json(args.actual_config_json)
    for row in runs:
        if row.get('config_path'):
            path = e.register(row['config_path'])
            if row.get('config_sha256') and hashlib.sha256(path.read_bytes()).hexdigest() != row['config_sha256']:
                raise ValueError('Saved run config changed after summary; regenerate summary')
    if not config:
        e.warnings.append('Summary lacks savedconfig_snapshot; regenerate with current summarizer or pass --actual-config-json. Missing settings remain unknown.')
    physics_meta_path = getattr(args, 'physics_capture_meta', None)
    if physics_meta_path is None and primary:
        candidate = e.path(summary['root']) / primary['episode_directory'] / 'physics_capture_meta.json'
        if candidate.is_file():
            physics_meta_path = candidate
    physics_meta = e.json(physics_meta_path) if physics_meta_path else {}
    builder_path = e.register('demos/lab_cuvette_demo04/run.py', required=False)
    backend = backend_snapshot(physics_meta, config, builder_path.read_text() if builder_path else '')
    frozen = e.json(args.frozen_scene) if args.frozen_scene else {}
    if frozen and get(quota, 'qualification_evidence.frozen_scene_sha256'):
        if frozen.get('frozen_scene_sha256') != get(quota, 'qualification_evidence.frozen_scene_sha256'):
            raise ValueError('Report frozen-scene evidence differs from qualified collection')
    portable = e.json(args.portable_validation) if args.portable_validation else {}
    process = e.json(args.process_evidence) if args.process_evidence else {}
    media, media_manifest = _media_items(e, args.media_manifest)
    figures = [item for item in media if item['kind'] == 'image']
    for filename, label, role in (
        ('control_vs_1000_visual_collision_overlay.png', '原12.5mm对照与10mm假设资产的实际视觉和28碰撞输入叠加', 'overlay'),
        ('recommended_1000_original_A_path.png', '10mm器皿与原A座入口及连续竖直扫掠的全部XY可行域', 'geometry')):
        path = e.register(geometry_path / filename)
        figures.append({'path': str(path), 'label': label, 'caption': label, 'kind': 'image', 'role': role})
    evidence_files = []
    for filename in ('candidate_results.csv', 'actual_runs.csv', 'standard5_runs.csv', 'collection_status.csv'):
        p = e.register(summary_path.parent / filename, required=False)
        if p:
            evidence_files.append({'path': str(p), 'label': filename})
    for filename in ('candidate_dimensions_and_geometry.csv', 'dimensions_mass_and_inertia.csv', 'asset_derivation_checks.json'):
        p = e.register(geometry_path / filename)
        evidence_files.append({'path': str(p), 'label': filename})
    if audit_path:
        evidence_files.append({'path': str(e.register(audit_path)), 'label': '数据资格与正常读取验证'})
    if args.frozen_scene:
        evidence_files.append({'path': str(e.register(args.frozen_scene)), 'label': '冻结standard场景清单'})
    if failure1050:
        evidence_files.append({'path': str(e.register(FAILURE1050)), 'label': '10.50mm开发失败复核'})
    try:
        head = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=e.root, text=True).strip()
        branch = subprocess.check_output(['git', 'branch', '--show-current'], cwd=e.root, text=True).strip()
    except (OSError, subprocess.CalledProcessError):
        head = branch = 'unknown'
    return {'schema': 'labsim.Size_v1.report_evidence.v1',
        'created_at': datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),
        'root': str(e.root), 'summary': summary, 'geometry': geometry, 'registry': registry,
        'quota': quota, 'collection_audit': audit, 'primary': primary,
        'primary_result': result, 'actual_config': config, 'frozen_scene': frozen,
        'portable_validation': portable, 'process_evidence': process,
        'cooked': cooked, 'failure1050_audit': failure1050,
        'physics_capture_meta': physics_meta, 'physics_backend': backend,
        'media_manifest': media_manifest, 'media': [x for x in media if x['kind'] == 'video'],
        'figures': figures, 'evidence_files': evidence_files, 'sources': e.sources,
        'warnings': e.warnings, 'authoring_head': head, 'branch': branch,
        'authoring_runtime': BUNDLED_PYTHON,
        'render_required': 'Packaged render_docx.py and bundled LibreOffice only; inspect every rendered page'}


def table(header, rows, widths):
    return ('table', {'header': header, 'rows': rows, 'widths_cm': widths})


def image(item, height=7):
    return ('image', {**item, 'caption': item.get('caption', item['label']), 'max_height_cm': height})


def pages_for(d):
    s, g, q, c = d['summary'], d['geometry'], d['quota'], d['actual_config']
    p, r = d['primary'] or {}, d['primary_result']
    asset_id = p.get('asset_id') or g['recommended_initial_physics_candidate']['asset_id']
    asset = d['registry']['assets'].get(asset_id, {})
    true_runs = [x for x in s.get('runs', []) if x.get('task_type') == 'loading'
        and x.get('task_success') is True and x.get('complete') is True
        and all(x.get(k) is True for k in ('release_observed', 'retreat_observed', 'stable_after_retreat'))]
    full = bool(true_runs)
    title = 'LabSim 比色皿尺寸变体装载实验报告'
    lead = (f"当前实际独立评分已确认完整装载，成功资产为{', '.join(sorted({x['asset_id'] for x in true_runs}))}。" if full else
            '当前没有已确认的完整任务成功。已完成尺寸资产与几何筛选，实际装载结果及中止阶段见本报告。')
    pages = []
    pages.append([('title', title), ('p', 'Size_v1  2026年10月21日  固定原A座与standard场景'),
        ('p', lead + '本轮只改变器皿XY尺寸，原座、整机、来源架和物理参数保持。物理任务判分、进程退出和正式冻结数据资格分别记录；开发成功不自动增加配额。尺寸是实验假设，不是实物反推或平台精度认证。'),
        ('h', '本轮结果'),
        table(['项目', '实际状态', '证据范围'], [
            ['尺寸资产', f"{g['evaluated_count']}个已生成并评估", '9.50至12.50mm；独立USD与asset_id全部保留'],
            ['原座', 'candidate_far_hinge 原A座', '保持原始几何和内部结构；未扩大孔口或删除凸起'],
            ['完整任务', '已确认成功' if full else '尚未确认成功', '实际释放 撤离 稳定必须同时成立'],
            ['5次standard', get(s, 'standard5.status', '未知'), f"完成记录{get(s, 'standard5.closed_record_count', '未知')}条；逐次结果单列"],
            ['合格成功', num(q.get('qualified_success_count'), 0) + ' / 20', '需冻结场景与正常数据读取资格'],
            ['合格失败', num(q.get('qualified_failure_count'), 0) + ' / 20', '自然失败与动作偏差失败分组'],
            ['random及训练', '本轮未要求random 不新增训练', '训练和自动验证默认固定standard']
        ], [3.2, 4.0, 9.6]),
        ('p', '本报告供项目协作者和后续agent复核：哪些内容真正运行、使用哪个尺寸、剩余配额和问题是什么。几何筛选、物理验证、完整任务成功和数据资格分别记录，未知值保留。'),
        ('p', f"开发分支 {d['branch']}。报告生成时HEAD {d['authoring_head']}。实际episode另保存自己的代码和控制版本；不把报告提交误当作运行提交。")])
    scene_figure = next((x for x in d['figures'] if x.get('role') in {'overview', 'scene', 'main_frame', 'dual_view'}),
                       next((x for x in d['figures'] if x.get('role') not in {'overlay', 'geometry'}), None))
    if scene_figure:
        pages[0].append(image(scene_figure, 5.2))

    overlay = next(x for x in d['figures'] if x.get('role') == 'overlay')
    rows = []
    for ident in ('cuvette_xy_1250', asset_id):
        a = d['registry']['assets'][ident]
        rows.append([ident, ' × '.join(num(x, 2) for x in a['nominal_outer_xy_mm']),
            num(a['actual_outer_dimensions_mm'][0], 6) + ' × ' + num(a['actual_outer_dimensions_mm'][1], 6),
            num(a['height_nominal_mm'], 2), num(a['mass_kg'] * 1000, 1)])
    pages.append([('h', '资产派生与尺寸依据'),
        ('p', f"从保留的12.5mm控制资产按diag(w/12.5,w/12.5,1)派生。{g['evaluated_count']}个候选采用同一变换，包括视觉和28块厚壁碰撞输入；空腔、底部、倒角拓扑和底面原点不变，不制作局部避让。每个参与评估的尺寸均先注册并重载核查。"),
        table(['资产ID', '名义XY mm', '实际USD外XY mm', '高度 mm', '假设质量 g'], rows, [4.1, 2.6, 5.5, 2.2, 2.4]),
        ('p', '实际顶点包络中的微小舍入差用于检查文件一致性，不能称为实测尺寸精度。其他架内同型号器皿引用同一选定资产；来源架仍为完整6×6参数化架，16mm槽位和孔距不随器皿缩小。'),
        image(overlay, 9),
        ('p', '总质量沿用20g工程假设。均匀材料体积分包含底部和杯壁、排除空腔，重新计算COM和完整惯量；总质量固定，因此数值密度随尺寸改变，不同时假设材料密度固定。没有把重叠凸体体积相加作为质量。'),
        ('p', f"源控制资产SHA256 {g['source_control_sha256']}。原仪器资产SHA256 {g['original_instrument_sha256']}。全35项内腔截面、COM、惯量、接触配置与哈希在同包CSV及注册清单中。")])

    candidate_map = {x['asset_id']: x for x in s.get('candidates', [])}
    rows = []
    for ident in ('cuvette_xy_1250', 'cuvette_xy_1102', 'cuvette_xy_1100', 'cuvette_xy_1072', 'cuvette_xy_1070', 'cuvette_xy_1050', 'cuvette_xy_1030', 'cuvette_xy_1000', 'cuvette_xy_0950'):
        a = candidate_map.get(ident)
        if a:
            rows.append([num(a['nominal_outer_width_mm'], 2), state(a['entry_required_slice_four_yaw_XY_domain_present']),
                '正证书' if a['conservative_full_path_certificate_present'] else '未认证',
                num(a.get('conservative_XY_clearance_lower_bound_mm'), 4),
                short(a.get('complete_task_result'), 24)])
    geometry_figure = next(x for x in d['figures'] if x.get('role') == 'geometry')
    pages.append([('h', '原座通道与可行尺寸范围'),
        table(['宽度 mm', '入口全XY域', '全深竖直路径', '保守余量 mm', '自然控制完整任务'], rows, [2.0, 2.8, 3.4, 3.0, 5.6]),
        ('p', '尺寸表的完整任务结果仅汇总明确标注condition_group=natural的未扰动控制组。确定性动作偏差的成败另列，不将受控提前停止等失败解释为该尺寸自然不稳定；未知组不归入natural。全部原始成败仍保留在候选CSV和运行列表，正式配额继续按各条件组独立审计。'),
        ('p', '10.70mm是当前采样中取得保守全深正证书的最大宽度，不是全局物理尺寸上限。10.72至11.00mm在入口有可行域，但保守全深检查未认证，不能写成必然无法装载。11.02mm起的已测试尺寸在四个声明直立方位的入口必要域为空；不排除未研究的倾斜6DoF路线。'),
        image(geometry_figure, 7.5),
        ('p', '所有尺寸共用0 90 180 270度方位和全部有界XY平移域。入口取真实座高28.5mm闭合孔域与器皿底上1mm实心截面，计算三角形对Minkowski禁区并集。全深正证书投影保留原座所有承托面上0.1µm以上三角面和垂直壁线段，并用真实器皿外轮廓作保守连续竖直扫掠。负的保守检查仅表示未证明。'),
        ('p', '推荐初始物理候选为10.00mm，原A座局部偏移X=+0.650mm Y=+0.001mm，保守侧向输入表面余量约0.350mm。原座约2mm承托高度与约27mm插入深度是网格几何候选；最低底面承托采样凸包包含COM，实际释放稳定仍须仿真验证。')])

    runs = [x for x in s.get('runs', []) if x.get('task_type') == 'loading']
    run_table = [[short(x.get('episode_id') or x['run_name'], 12), short(x.get('asset_id'), 21), num(x.get('simulation_seconds'), 2),
        state(x.get('task_success')), short(x.get('termination_reason'), 43)] for x in runs[-7:]]
    if not run_table:
        run_table = [['无完成装载记录', '未知', '未知', '未执行或未闭合', '以actual_runs.csv为准']]
    phase = [('抓取', r.get('grasp_success')), ('实际取出', r.get('source_slot_cleared')),
        ('搬运', r.get('transport_success')), ('装载序列', r.get('loading_sequence_observed')),
        ('实际松爪', r.get('release_observed')), ('实际撤离', r.get('retreat_observed')),
        ('撤离后稳定', r.get('stable_after_retreat'))]
    failure_blocks = [('p', paragraph) for paragraph in failure1050_paragraphs(d.get('failure1050_audit', {}))]
    failure_link = next((item for item in d['evidence_files'] if item['label'] == '10.50mm开发失败复核'), None)
    pages.append([('h', '实际连续任务与失败记录'),
        ('p', '控制沿用连续接触锁存：首次双指接触后保持已有闭合命令，避免锁存切换时重新张开。新宽度由实际接触适应，不复制12.5mm的固定指距。示范控制器使用器皿真值位姿；本轮未训练模型。'),
        table(['episode ID', '资产', '仿真秒', '独立结果', '终止原因'], run_table, [3.1, 3.8, 1.8, 2.3, 5.8]),
        ('p', f"主记录 {p.get('episode_directory', '尚无可选的已闭合装载episode')}。任务结果只取独立评分，不根据INSERT或VERIFY阶段名称认定成功。完整视频必须来自同一次运行，不拼接不同尝试。"),
        ('p', '主记录独立观察：' + '；'.join(label + state(value) for label, value in phase) + '。'),
        ('p', f"当前主记录终止原因 {r.get('termination_reason', '未知')}；失败类型 {r.get('failure_type', '未知')}。物理质量标记 {', '.join(r.get('physics_quality_flags', [])) or '无已记录项或尚无完成评分'}。完整已启动及未闭合列表另存CSV。"),
        ('h', '十点五毫米开发运行的独立失败'),
        *failure_blocks,
        *([('link', failure_link)] if failure_link else [])])

    verification = get(s, 'standard5', {})
    selected = set(verification.get('episodes', []))
    rows = [[short(x['episode_directory'], 37), state(x.get('task_success')),
             num(x.get('simulation_seconds'), 2), state(x.get('record_valid')), short(x.get('termination_reason'), 38)]
            for x in s.get('runs', []) if x['episode_directory'] in selected]
    if not rows:
        rows = [['尚未指定或尚未启动5次验证', '未知', '未知', '未知', '不以开发preview替代']]
    window = r.get('post_retreat_window', {})
    pages.append([('h', '固定standard重复性与稳定验收'),
        table(['episode', '任务结果', '仿真秒', '记录有效', '终止原因'], rows[:8], [5.2, 2.2, 1.7, 2.4, 5.3]),
        ('p', f"显式选入{verification.get('selected_episode_count', '未知')}条，完成记录{verification.get('closed_record_count', '未知')}条，独立成功{verification.get('independently_scored_success_count', '未知')}条，独立失败{verification.get('independently_scored_failure_count', '未知')}条。结果未知{verification.get('outcome_unknown_count', '未知')}条。5次验证中满足最终冻结配置和数据资格的成功可以计入20条。"),
        ('h', '释放与撤离的固定判据'),
        ('p', '接触后停止向下推进，建立底部承托，逐步卸载夹持，对称慢松爪，确认双指脱离，再撤离并观察至少3秒。位置漂移200µm、姿态变化1度及声明的50µm输入材料穿透门槛保持。原生速度与位姿差分都记录；信号不一致保留告警，不将其删除或据阶段直接写成功。'),
        table(['主记录撤离后窗口', '实际值', '判读'], [
            ['窗口状态', window.get('status', 'unavailable'), '撤离前中止时必须unavailable'],
            ['稳定', state(window.get('stable')), '不拿terminal_window代替'],
            ['测得时长 秒', num(window.get('measured_duration_s'), 3), '要求至少3秒'],
            ['位置漂移 mm', num(None if window.get('position_drift_m') is None else window['position_drift_m'] * 1000, 6), '本轮工程门槛0.2mm'],
            ['姿态变化 度', num(window.get('orientation_change_deg'), 6), '本轮工程门槛1度']
        ], [5.0, 3.5, 8.3]),
        ('p', '输入材料检查考虑实际器皿姿态与原座网格，不只检查中心。它与实际接触力、接触候选和PhysX烹饪表示不同；50µm工程门槛不代表真实标定或适用于其他紧配插入任务。')])

    cooked = d['cooked']
    convex = cooked.get('convex', {})
    side = get(c, 'cameras.side_camera', {})
    wrist = get(c, 'cameras.wrist_camera', {})
    backend = d.get('physics_backend', {})
    if backend.get('backend_source') == 'shared_builder_code_declaration':
        backend_label = 'PhysX CPU dynamics 共享构建代码声明'
    else:
        gpu = backend.get('gpu_dynamics')
        backend_label = str(backend.get('backend') or '未知') + (' GPU dynamics' if gpu is True else ' CPU dynamics' if gpu is False else ' dynamics未知')
    ccd_label = '开启' if backend.get('ccd') is True else '关闭' if backend.get('ccd') is False else '未知'
    pages.append([('h', '物理配置与烹饪表示的证据边界'),
        table(['参数', '实际配置或证据', '范围'], [
            ['后端与CCD', backend_label + '；CCD ' + ccd_label, '后端来源' + str(backend.get('backend_source', 'unknown')) + '；CCD来源' + str(backend.get('ccd_source', 'unknown'))],
            ['步进与控制', f"{get(c, 'physics.hz', '未知')}Hz / {get(c, 'physics.control_hz', '未知')}Hz", '每控制周期物理步数由频率比计算'],
            ['求解器', str(get(c, 'physics.solver', '未知')), f"位置迭代{get(c, 'physics.position_iterations', '未知')} 速度迭代{get(c, 'physics.velocity_iterations', '未知')}"],
            ['接触偏移与静止偏移', f"{num(get(c, 'physics.contact_offset_m'), 8)}m / {num(get(c, 'physics.rest_offset_m'), 8)}m", '沿用基线；不将contact_offset称为位置精度'],
            ['器皿摩擦与恢复', f"静{get(c, 'physics.static_friction', '未知')} 动{get(c, 'physics.dynamic_friction', '未知')} 恢复{get(c, 'physics.restitution', '未知')}", '假设值与average组合；未真实校准'],
            ['夹持力限制', f"{get(c, 'env.gripper_effort_max_N', '未知')}N", '按阶段调度记录；不针对每个尺寸加力或调摩擦']
        ], [4.0, 5.7, 7.1]),
        ('p', f"preview探针中28个碰撞输入均取得有效凸体烹饪返回。世界米坐标重算的返回表面到输入表面抽样最大{num(None if convex.get('max_observed_cooked_to_input_surface_distance_m') is None else convex['max_observed_cooked_to_input_surface_distance_m'] * 1e6, 3)}µm；反向为{num(None if convex.get('max_observed_input_to_cooked_surface_distance_m') is None else convex['max_observed_input_to_cooked_surface_distance_m'] * 1e6, 3)}µm。零顶点误差不代表整个烹饪面无误差。"),
        ('p', '这些是逐块顶点和面中心到三角表面的已观察距离，不是复合几何并集外表面误差或全局Hausdorff上界。接口返回当前prim的烹饪表示，未读取已附着PxShape指针。原座实际物理raycast命中表面与源三角面一致的抽样证据，不能替代完整静态烹饪顶点索引读取，也不排除未采样差异。'),
        ('h', '连续相机与时间对齐'),
        ('p', f"侧相机固定世界安装，腕相机安装于panda_hand。实际主配置侧视分辨率{side.get('resolution', '未知')}，焦距{side.get('focal_length_cm', '未知')}cm；腕视{wrist.get('resolution', '未知')}，焦距{wrist.get('focal_length_cm', '未知')}cm。镜头安装、裁剪与策略256像素变换保持显式配置。"),
        ('p', '双相机在完整任务过程中持续采集；同步记录观测 动作 后继状态和仿真时间。renderer当前状态对齐的核验与真实相机曝光或实物标定分开。实际本轮未完成的动态同步核验不得用旧preview结果代替；详情按同包原始核验JSON判读。')])

    grouped = q.get('condition_groups', {})
    reader = d['collection_audit'].get('reader_validation', [])
    pages.append([('h', '冻结场景与训练数据资格'),
        ('p', 'collect train以及训练内自动验证统一使用standard并关闭布局 资产 光照和物理随机化。训练RNG种子另行配置；设置随机种子不等于关闭场景随机化。本轮取消random10，不把它列为剩余验收。'),
        table(['数据组', '合格成功', '合格失败'], [
            ['总计 目标各20条', num(q.get('qualified_success_count'), 0), num(q.get('qualified_failure_count'), 0)],
            ['自然控制', num(get(grouped, 'natural.qualified_success'), 0), num(get(grouped, 'natural.qualified_failure'), 0)],
            ['确定性动作偏差', num(get(grouped, 'deterministic_action_deviation.qualified_success'), 0), num(get(grouped, 'deterministic_action_deviation.qualified_failure'), 0)],
            ['配额缺口', num(q.get('success_missing'), 0), num(q.get('failure_missing'), 0)]
        ], [7.0, 4.9, 4.9]),
        ('p', f"资格状态 {q.get('status', '未知')}。几何预检 资产加载 初始化 服务异常 人工中断及未真正执行任务的preview不计失败配额。确定性动作偏差只改变已声明阶段动作，记录幅度与版本，不改变冻结场景，不强撞或伪造标签。"),
        ('p', '成功BC默认只读record_valid且task_success的合格完整loading轨迹；自然失败和受控偏差失败分别保存。完整episode划分训练和验证集合，两者保持同一冻结场景；相邻帧不跨集合。原始失败标签不为读取器修改。'),
        table(['读取与冻结项', '实际证据', '说明'], [
            ['正常读取验证条目', str(len(reader)), '每项由已有ACT或DP正常入口实际执行'],
            ['冻结场景指纹', short(d['frozen_scene'].get('frozen_scene_sha256') or get(q, 'qualification_evidence.frozen_scene_sha256'), 36), '完整指纹与资产依赖在JSON'],
            ['便携episode验证', '已提供证据' if d['portable_validation'] else '未提供本轮证据', '不把旧绝对路径外链当成便携副本'],
            ['标签与资产', '保持原始值', 'asset_id 场景 控制与代码版本随episode保存']
        ], [4.0, 5.2, 7.6]),
        ('p', '资格计数未知时保留未知，不把它自动填成0或20。采集审计与汇总不一致时应先重新核查，不开始训练。训练模型不是本轮执行内容。')])

    video_blocks = [('media', item) for item in d['media'][:5]]
    evidence_blocks = [('link', item) for item in d['evidence_files'][:4]]
    pages.append([('h', '复现与交付入口'),
        ('p', f"选定原A座，推荐或实际使用资产{asset_id}。实际完整任务结论为{state(r.get('task_success'))}，释放为{state(r.get('release_observed'))}，撤离为{state(r.get('retreat_observed'))}，撤离后稳定为{state(r.get('stable_after_retreat'))}。原12.5mm控制组合的既有冲突结论保留。"),
        ('code', 'bash demos/lab_cuvette_size_v1/run.sh --mode teacher \\\n  --config config/lab_cuvette_loading/Size_v1/standard.yaml \\\n  --asset-id ' + asset_id + ' --output NEW_EPISODE_DIRECTORY\n\n'
          'python3 scripts/summarize_lab_cuvette_size_v1.py \\\n  --root outputs/lab_cuvette_loading/Size_v1 \\\n  --output NEW_SUMMARY_DIRECTORY'),
        ('p', '默认GUI；服务器环境可加--headless。每次使用新目录。正式collect train和自动验证锁定冻结清单中的asset_id，不能随意改尺寸；尺寸对比须使用独立开发配置和输出，不计最终10mm固定场景配额。解析器保留当前目标偏移，不自动为其他尺寸调参或继承10mm结果。以上命令存在不表示5次或20+20已完成。'),
        ('h', '同包代表录像'),
        *video_blocks,
        ('p', '视频仅按manifest注明整段运行或剪辑范围；主录像应连续包含其实际执行阶段。未发生的插入 松爪 撤离不通过剪辑包装为成功。'),
        ('h', '结果表与剩余工作'),
        *evidence_blocks,
        ('p', '若完整任务仍未通过，优先依据准确失败阶段 真实接触与源几何重叠记录修复控制和时序；原座不修改。若已通过，先核对固定5次重复性，再按同一冻结场景补齐成功与有效失败配额。实际缺口按本报告数据表，不预定成功。所有参数仍待实物测量校正，real_calibrated=false。')])
    return pages


def relative_target(item, parent):
    path = Path(item['path']).resolve()
    if not path.is_relative_to(parent.resolve()):
        raise ValueError('Deliverable link must be copied inside the report folder')
    return path.relative_to(parent.resolve()).as_posix()


def write_markdown(pages, path):
    lines = []
    for page in pages:
        for kind, value in page:
            if kind in {'title', 'h'}:
                lines += [('## ' if kind == 'h' else '# ') + value, '']
            elif kind == 'table':
                lines += ['| ' + ' | '.join(value['header']) + ' |', '| ' + ' | '.join(['---'] * len(value['header'])) + ' |']
                lines += ['| ' + ' | '.join(str(x).replace('|', '/') for x in row) + ' |' for row in value['rows']]
                lines.append('')
            elif kind == 'image':
                lines += [f"![{value['caption']}]({relative_target(value, path.parent)})", value['caption'], '']
            elif kind in {'media', 'link'}:
                lines += [f"[{value['label']}]({relative_target(value, path.parent)})", '']
            elif kind == 'code':
                lines += ['```bash', value, '```', '']
            else:
                lines += [value, '']
    path.write_text('\n'.join(lines), encoding='utf-8')


def write_docx(pages, path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from PIL import Image
    doc = Document()
    for style in doc.styles:
        for border in list(style.element.findall('.//' + qn('w:pBdr'))):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in (('Normal', 10.0), ('Title', 20), ('Heading 1', 14), ('Caption', 8.8)):
        style = doc.styles[name]
        style.font.name = 'Aptos'; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
        style.paragraph_format.line_spacing = 1.12; style.paragraph_format.space_after = Pt(5)
    doc.styles['Title'].paragraph_format.space_after = Pt(11)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(7)
    for index, page in enumerate(pages):
        if index:
            doc.add_page_break()
        for kind, value in page:
            if kind in {'title', 'h'}:
                doc.add_paragraph(value, 'Title' if kind == 'title' else 'Heading 1')
            elif kind == 'table':
                table_obj = doc.add_table(rows=1, cols=len(value['header']))
                table_obj.alignment = WD_TABLE_ALIGNMENT.CENTER; table_obj.autofit = False
                for j, text in enumerate(value['header']):
                    table_obj.rows[0].cells[j].text = text
                for row in value['rows']:
                    cells = table_obj.add_row().cells
                    for j, text in enumerate(row):
                        cells[j].text = str(text)
                repeat = OxmlElement('w:tblHeader'); table_obj.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row_index, row in enumerate(table_obj.rows):
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for j, cell in enumerate(row.cells):
                        cell.width = Cm(value['widths_cm'][j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        prop = cell._tc.get_or_add_tcPr()
                        borders = OxmlElement('w:tcBorders')
                        for edge in ('top', 'left', 'bottom', 'right'):
                            border = OxmlElement('w:' + edge)
                            for k, v in {'val': 'single', 'sz': '5', 'color': 'D9D9D9'}.items():
                                border.set(qn('w:' + k), v)
                            borders.append(border)
                        prop.append(borders)
                        margins = OxmlElement('w:tcMar')
                        for edge in ('top', 'left', 'bottom', 'right'):
                            margin = OxmlElement('w:' + edge); margin.set(qn('w:w'), '75'); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
                        prop.append(margins)
                        if row_index == 0:
                            fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'E6EDF3'); prop.append(fill)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(0); paragraph.paragraph_format.line_spacing = 1.08
                            if j in (1, 2) and len(value['header']) > 3:
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in paragraph.runs:
                                run.font.size = Pt(9.0); run.bold = row_index == 0
                doc.add_paragraph('').paragraph_format.space_after = Pt(0)
            elif kind == 'image':
                with Image.open(value['path']) as picture:
                    w, h = picture.size
                width = min(16.5, float(value.get('max_height_cm', 7)) * w / h)
                paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                paragraph.add_run().add_picture(value['path'], width=Cm(width))
                doc.add_paragraph(value['caption'], 'Caption')
            elif kind in {'media', 'link'}:
                paragraph = doc.add_paragraph(); link = OxmlElement('w:hyperlink')
                relation = paragraph.part.relate_to(relative_target(value, path.parent), RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                link.set(qn('r:id'), relation)
                run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = value['label']
                run.append(text); link.append(run); paragraph._p.append(link)
            elif kind == 'code':
                paragraph = doc.add_paragraph(value); paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.name = 'DejaVu Sans Mono'; run.font.size = Pt(7.6)
            else:
                doc.add_paragraph(value)
    doc.core_properties.title = 'LabSim 比色皿尺寸变体装载实验报告'
    doc.core_properties.subject = 'Size_v1 原座固定尺寸假设 实际任务与数据资格'
    doc.core_properties.author = 'LabSim 项目'
    doc.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=ROOT)
    parser.add_argument('--summary', type=Path, required=True)
    parser.add_argument('--geometry', type=Path)
    parser.add_argument('--collection-audit', type=Path)
    parser.add_argument('--media-manifest', type=Path)
    parser.add_argument('--frozen-scene', type=Path)
    parser.add_argument('--portable-validation', type=Path)
    parser.add_argument('--process-evidence', type=Path)
    parser.add_argument('--physics-capture-meta', type=Path, help='Optional actual physics metadata JSON; defaults to selected episode metadata when present')
    parser.add_argument('--actual-config-json', type=Path, help='Optional existing exact run configuration JSON; YAML parsing belongs to the existing summary data tool')
    parser.add_argument('--primary-episode-id')
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--stem', default='LabSim_Size_v1_尺寸变体装载实验报告')
    parser.add_argument('--media-budget-mib', type=float, default=30)
    parser.add_argument('--validate-only', action='store_true')
    args = parser.parse_args(argv)
    if args.output.exists():
        raise FileExistsError('Use a NEW report directory')
    e = Evidence(args.root)
    data = collect(e, args)
    if args.validate_only:
        pages = pages_for(data)
        print(json.dumps({'validated_without_authoring': True, 'pages_planned': len(pages),
            'runs': len(data['summary'].get('runs', [])), 'videos': len(data['media']),
            'primary_task_success': data['primary_result'].get('task_success'),
            'qualified_success_count': data['quota'].get('qualified_success_count'),
            'qualified_failure_count': data['quota'].get('qualified_failure_count')}, ensure_ascii=False))
        return data
    budget = sum(Path(x['path']).stat().st_size for x in data['figures'] + data['media'])
    if budget > args.media_budget_mib * 1024**2:
        raise ValueError('Selected media exceeds compact report budget; provide smaller representative media or explicitly raise budget')
    args.output.mkdir(parents=True, exist_ok=False)
    for folder, items in (('media', data['figures'] + data['media']), ('data', data['evidence_files'])):
        destination = args.output / folder; destination.mkdir()
        for index, item in enumerate(items):
            source = Path(item['path'])
            target = destination / f'{index:02d}_{source.name}'
            shutil.copy2(source, target); item['path'] = str(target.resolve())
    pages = pages_for(data)
    write_markdown(pages, args.output / (args.stem + '.md'))
    write_docx(pages, args.output / (args.stem + '.docx'))
    (args.output / 'report_evidence.json').write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')
    print(json.dumps({'docx': str(args.output / (args.stem + '.docx')), 'planned_pages': len(pages),
        'render_and_all_page_visual_QA_required': True, 'videos': len(data['media'])}, ensure_ascii=False))
    return data


if __name__ == '__main__':
    main()

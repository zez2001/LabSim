"""Build the concise internal-use feedback with workspace-bundled python-docx."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
def text(d,s,style=None,size=10.5):
 p=d.add_paragraph(style=style);p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=Pt(15)
 r=p.add_run(s);r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0);r.font.name='Arial';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');return p
def heading(d,s):return text(d,s,'Heading 1',14)
def code(d,s):
 for line in s.splitlines():
  p=text(d,line,size=9);p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=Pt(12)
def main():
 summary=json.loads((ROOT/'teacher_closed/summary.json').read_text());reader=json.loads((ROOT/'teacher_closed/read_success.json').read_text())
 d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=sec.bottom_margin=Cm(1.9);sec.left_margin=sec.right_margin=Cm(2)
 for n in ['Normal','Title','Heading 1','Heading 2']:
  st=d.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 text(d,'LabSim 内部使用优化反馈','Title',20)
 text(d,'2026年10月28日  Project_release_v1  codex/labsim',size=10)
 text(d,'已统一当前文档与完整Teacher入口，增加数据用途选择和移交检查，修正图像输入契约并完成实际tokenizer审计。普通Teacher完整装载已实际通过，关闭后由正式ACT和DP入口读出成功batch。场景、原座、器皿尺寸、控制与评分沿用发布基线。')
 heading(d,'使用入口与修改范围')
 text(d,'从仓库README开始，随后按docs/development/ONBOARDING.md执行环境准备、完整任务、标准采集、读取和策略接入。当前版本快捷别名current解析为Project_release_v1；省略版本继续兼容Project_v1。启动摘要集中显示真实版本、任务、策略、资产、指纹、模式、用途及输出目录。')
 text(d,'用途record_purpose独立于成功资格。demonstration用于示范，policy_evaluation用于策略，lifecycle_diagnostic与resource_probe用于诊断。历史用途通过episode身份、来源路径和完整SHA256绑定sidecar补充，原H5和任务标签不改。')
 heading(d,'本次实际验证')
 rows=[('检查','实际结果'),('CPU与文档','904项CPU通过；138本地链接和19条命令检查通过'),('普通Teacher',f"完整成功；{summary['transitions']}条transition；3秒撤离后稳定"),('记录与读取','版本、资产、用途与指纹闭合；正式ACT和DP成功batch可读'),('历史保护','14个原H5及任务标签哈希前后相同；生命周期资格2／1／3'),('用途集合','π0.5正式策略1条；显式诊断11条；资源探针10条'),('数据移交','relative-links整树移动、copy独立移动可读；缺源文件明确报告')]
 t=d.add_table(rows=0, cols=2);t.autofit=False;t.columns[0].width=Cm(3.4);t.columns[1].width=Cm(13.4)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,v in enumerate(row):
   cells[j].text=v
   for p in cells[j].paragraphs:
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
    for r in p.runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor(0,0,0);r.font.name='Arial';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
   tcpr=cells[j]._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for n in ['top','left','bottom','right']:
    e=OxmlElement('w:'+n);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tcpr.append(b)
   if i==0:
    e=OxmlElement('w:shd');e.set(qn('w:fill'),'D9E2F3');tcpr.append(e)
    for r in cells[j].paragraphs[0].runs:r.bold=True
 p=text(d,'完整原始连续录像见随包media/teacher_full.mp4；H5约2 GB，按来源清单单独移交。所有输出使用新目录，历史结果保持原样。',size=10);p.paragraph_format.space_before=Pt(7)
 d.add_page_break()
 heading(d,'推荐调用顺序')
 text(d,'先选择已安装Isaac Sim 5.1的Python 3.11。完整命令见随包commands.sh；以下均从仓库根目录执行。')
 code(d,'./labsim assets download\n./labsim doctor --project-version Project_release_v1 --json\n./labsim demo --project-version Project_release_v1 --backend teacher \\\n  --gui --record-mode collect --record-purpose demonstration \\\n  --output outputs/demo_NEW')
 text(d,'标准示范采集与成功读取')
 code(d,'export LABSIM_RUN=outputs/collect_NEW\n./labsim collect --project-version Project_release_v1 --backend teacher \\\n  --headless --episodes 1 --record-purpose demonstration \\\n  --record-mode collect --output "$LABSIM_RUN"\n# 等待进程关闭\n"$LABSIM_PYTHON" -m demos.lab_cuvette_project_release_v1.dataset_index \\\n  "$LABSIM_RUN" --purpose demonstration \\\n  --output "$LABSIM_RUN/index/demonstrations.hdf5"\n./labsim dataset-check --index "$LABSIM_RUN/index/demonstrations.hdf5" \\\n  --sample-filter success --policy both --output "$LABSIM_RUN/read.json"')
 text(d,'策略环境服务')
 code(d,'./labsim serve --project-version Project_release_v1 --headless \\\n  --record-purpose policy_evaluation --record-mode eval --port 8097 \\\n  --output outputs/rpc_NEW')
 text(d,'公开生命周期为reset、observe、execute、receipt、status、finish、close。模型等待时物理暂停，动作和后继状态分别记录。Teacher完整装载；GPT图像闭环抓取离架搬运；π0.5真实连续RPC；Harness本地短程接入；ACT／DP正常reader；自有策略使用同一接口。')
 heading(d,'输入事实与后续事项')
 text(d,'环境及π0.5服务使用原生640×480 HWC RGB uint8，线上带batch为[1,480,640,3]。H5／训练reader分支为256×256 CHW；checkpoint另用PIL双线性缩至224×168、上下填边，再转模型浮点张量。现有图像及动作转换数值保持。')
 text(d,'实际部署tokenizer测得任务文本264→200，截断在TCP四元数中。保存文本哈希、完整及有效tokens、mask与decode；历史内部张量为同源码离线重建，新服务会直接记录processor实际输出。后续文本优化应使用独立输入配置。')
 text(d,'内部使用影响排序：先处理任务文本截断；大批次采集监控RSS并分组正常关闭；SDK退出回调保留日志和完成凭据；跨机器离线导入及PID身份改造单列后续。精细控制结果测量对象为空夹爪仿真TCP。资产质量摩擦及接触仍采用工程假设，真实校准单独进行。')
 text(d,'实际GPU运行提交 5c6b6422f9ba6fc16ac7b4374799c5ed7e426e47。最终发布提交、对应CI和逐文件SHA256见阅读包清单及publication.json。',size=9)
 # Replace inherited nil borders rather than appending competing definitions.
 for table in d.tables:
  pr=table._tbl.tblPr
  for old in list(pr.findall(qn('w:tblBorders'))):pr.remove(old)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'D9D9D9');borders.append(el)
  pr.insert(1,borders)
  for row in table.rows:
   for cell in row.cells:
    pr=cell._tc.get_or_add_tcPr()
    for old in list(pr.findall(qn('w:tcBorders'))):pr.remove(old)
 # Remove bundled template paragraph borders, including the Title blue rule.
 for part in (d._element,d.styles._element):
  for border in list(part.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.core_properties.title='LabSim 内部使用优化反馈';d.core_properties.author='LabSim';d.core_properties.subject='当前使用入口与实际验收'
 d.save(ROOT/'LabSim_内部使用优化反馈_20261028.docx')
if __name__=='__main__':main()

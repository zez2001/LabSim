"""Build the four-page group-meeting DOCX; source videos are copied unchanged.

Run with the workspace-bundled Python. The authoring marker must be run once
before this script is first invoked. QA JSON and screenshots stay outside the
meeting package. Overrides may update current-round acceptance after it closes.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[4]
HERE = Path(__file__).resolve().parent
BASE = ROOT / "outputs/lab_cuvette_loading/Project_v1_1"
NAME = "LabSim_环境与评测进展_20261028.docx"


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for b in iter(lambda: f.read(4 * 1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


def font(run, size=11, heading=False, bold=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    pr = run._element.get_or_add_rPr()
    fonts = pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        pr.insert(0, fonts)
    for k in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + k), "Arial")
    fonts.set(qn("w:eastAsia"), "黑体" if heading else "宋体")
    return run


def para(doc, text="", *, size=11, bold=False, heading=False, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text), size, heading, bold)
    return p


def title(doc, text):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(10)
    font(p.add_run(text), 20, True, True)
    return p


def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(6)
    font(p.add_run(text), 14, True, True)
    return p


def link(doc, label, target):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    rel = p.part.relate_to(target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for k in ("ascii", "hAnsi", "cs"):
        rf.set(qn("w:" + k), "Arial")
    rf.set(qn("w:eastAsia"), "宋体")
    pr.append(rf)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22"); pr.append(sz)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "000000"); pr.append(col)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); pr.append(u)
    r.append(pr)
    t = OxmlElement("w:t"); t.text = label; r.append(t)
    h.append(r); p._p.append(h)
    return p


def picture(doc, filename, caption, width=16.9):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(HERE / filename), width=Cm(width))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = para(doc, caption, size=9, after=7)
    c.paragraph_format.keep_with_next = False
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, width in enumerate(widths):
        t.columns[j].width = Cm(width)
    for row in [t.rows[0]]:
        trpr = row._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader"); trpr.append(repeat)
    for j, text in enumerate(headers):
        t.rows[0].cells[j].text = text
    for row in rows:
        cells = t.add_row().cells
        for j, text in enumerate(row):
            cells[j].text = str(text)
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + edge)
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), "D9D9D9")
                borders.append(el)
            pr.append(borders)
            mar = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + edge)
                el.set(qn("w:w"), "75")
                el.set(qn("w:type"), "dxa")
                mar.append(el)
            pr.append(mar)
            if i == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7E7E7"); pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.16
                for r in p.runs:
                    font(r, 9.5, False, i == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--overrides", type=Path)
    ap.add_argument("--marker-already-run", action="store_true", required=True)
    a = ap.parse_args()
    cfg = json.loads(a.overrides.read_text()) if a.overrides else {}
    out = a.output.resolve(); out.mkdir(parents=True, exist_ok=True)
    videos = [
        (cfg.get("teacher_video", str(BASE / "teacher_p9_regression_01/episode_0000/dual_view.mp4")), "teacher_full.mp4"),
        (cfg.get("policy_video", str(BASE / "gpt_debug_08/episode_0000/dual_view.mp4")), cfg.get("policy_video_name", "gpt6_astra_incomplete.mp4")),
    ]
    media = []
    for source, dest in videos:
        source = Path(source).resolve(strict=True)
        target = out / dest
        if source != target:
            shutil.copyfile(source, target)
        assert sha(source) == sha(target)
        media.append({"source": str(source), "file": dest, "bytes": target.stat().st_size,
                      "sha256": sha(target), "transformation": "byte_for_byte_whole_video_copy"})

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0); section.right_margin = Cm(2.0)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption", "Hyperlink"):
        if name in doc.styles:
            s = doc.styles[name]
            s.font.name = "Arial"; s.font.color.rgb = RGBColor(0, 0, 0)
            pr = s.element.get_or_add_rPr()
            rf = pr.find(qn("w:rFonts"))
            if rf is None: rf = OxmlElement("w:rFonts"); pr.append(rf)
            for k in ("ascii", "hAnsi", "cs"): rf.set(qn("w:" + k), "Arial")
            rf.set(qn("w:eastAsia"), "黑体" if name in ("Title", "Heading 1", "Heading 2") else "宋体")
    # Word's inherited Title style may contain a blue bottom border.
    # A formal report uses plain black typography with no decorative rule.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(11)
    doc.core_properties.title = "LabSim 环境与评测进展"
    doc.core_properties.subject = "2026年10月28日组内汇报"
    doc.core_properties.author = "LabSim 项目"

    title(doc, "LabSim 环境与评测进展")
    para(doc, "2026年10月28日组会  |  lab_cuvette_loading", size=10, after=10)
    para(doc, cfg.get("opening", "LabSim 已具备可运行的器皿装载物理任务、完整专家示范、图像闭环策略接口和可读取的双视角轨迹。当前重点是把单次实验能力整理为可持续采集、可终止、可复现的项目环境；策略任务失败与环境故障分别判定。"))
    heading(doc, "环境现在能做什么")
    para(doc, "Franka 从完整白色6×6来源架的A1槽抓取动态比色皿，拔出并搬运到单台开盖仪器，进入原内部A座，建立承托后慢松爪、撤离，并观察至少3秒。运动依靠真实刚体接触，器皿没有绑定到机器人。")
    picture(doc, cfg.get("scene_picture", "teacher_scene.png"), "图1  固定standard场景  完整来源架 单台开盖仪器 双视角采集")
    para(doc, "冻结场景：比色皿10×10×67.5 mm，质量20 g为假设；A1取物，F6保留蓝色器皿；原A座未扩孔。物理240 Hz、控制60 Hz，场景随机化关闭。该尺寸是仿真假设资产，并非实物校准尺寸。", size=10)
    para(doc, cfg.get("baseline", "已闭合的Project_v1_1证据：固定配置teacher五条完整成功，来自五个独立Isaac进程；正常ACT/DP读取通过。它证明单集任务链与数据可用，不能替代本轮同进程连续reset验收。"), size=10)

    doc.add_page_break()
    heading(doc, "执行 采集与评测的验收状态")
    table(doc, ["能力", "实际状态", "可据此使用的范围"], cfg.get("capability_rows", [
        ["完整物理任务", "已验证", "teacher完成抓取至撤离后3秒稳定"],
        ["通用动作执行", "已修复并实测", "连续IK、实际回执、部分执行保护"],
        ["夹持与承托交接", "已验证", "接触保持、逐步卸载、实际双指脱离"],
        ["同进程连续三集", "本轮待验收", "需成功—受控失败—reset后成功"],
        ["采集与正常读取", "单集已验证", "双视角、状态动作、时间戳、标签"],
        ["多集索引与筛选", "本轮待验收", "success / failure / all_valid需核对"],
        ["GPT真实图像闭环", "已验证接入", "最新08稳定搬运后主动结束，未装载"],
        ["其他真实外部策略", "短程历史接入", "本轮连续生命周期尚未验收"],
        ["异常关闭与性能", "本轮待验收", "补终止闭合、资源与体积测量"],
    ]), [3.4, 3.4, 10.2])
    para(doc, cfg.get("changes", "已解决的核心执行问题包括IK分支跳变、限幅后轨迹失真，以及move/wait阶段重新张指导致滑移。新执行器从实际已下发状态连续规划，动作回执区分已拒绝、已完成、部分执行和未达到目标；夹持保持不会吸附器皿。支撑交接先停止顶压、逐步卸载，再实际张爪撤离。"))
    para(doc, "自由空间诊断中的100 μm微小位移属于仿真内TCP测量，不能作为接触装载精度或真机精度结论。原生速度与位姿差分差异、有效碰撞烹饪表示未完整读回等问题继续作为质量告警保存。", size=10)

    doc.add_page_break()
    heading(doc, "真实流程和零样本接入")
    picture(doc, cfg.get("retreat_picture", "teacher_retreat.png"), "图2  teacher实际张爪并撤离后  比色皿独立留在原内部A座")
    para(doc, cfg.get("teacher_details", "配套teacher视频来自2026年10月27日Project_v1_1 P9回归，运行代码9a88ea0。单次连续125.667秒仿真中，112.192秒观测到双指张开且脱离，116.417秒观测到机械臂实际撤离，随后完成3秒稳定检查。录像覆盖初态到结束，没有拼接或调速。"), size=10)
    picture(doc, cfg.get("policy_picture", "gpt08_terminal.png"), cfg.get("policy_caption", "图3  GPT08真实闭环的结束状态  已搬运 仍夹持 未完成装载"))
    para(doc, cfg.get("policy_details", "第二段视频为同日Project_v1_1的Cubence gpt-6-astra GPT08真实闭环。模型仅使用双图、相机标定、机器人自身状态和声明的指端接触信号；器皿及目标座真值不提供给策略。65.117秒仿真后模型主动finish：记录有效，任务未完成，未执行装载后的释放与撤离。此结果说明真实策略已接入，并非零样本装载成功。"), size=10)

    doc.add_page_break()
    heading(doc, "现在如何使用")
    para(doc, cfg.get("usage", "获取codex/labsim固定提交，选择已安装Isaac Sim 5.1的Python 3.11，下载固定revision资产并校验。通过./labsim显式选择Project_v1_1，分别运行demo、collect、dataset-check和policy接口。省略版本会进入旧Project_v1，首次使用应保留显式版本参数。"))
    para(doc, cfg.get("command", "./labsim demo --project-version Project_v1_1 --backend teacher \\\n  --headless --output outputs/lab_cuvette_loading/Project_v1_1/demo_NEW"), size=9.5, after=8)
    para(doc, "采集数据保存双相机图像、机器人状态、原始请求、实际动作、时间及独立标签。等待进程完全关闭后再运行正常读取器；成功BC只读取合格成功，任务失败、服务错误和未完成记录分别管理。模型请求期间仿真暂停，执行动作时按明确物理步推进。", size=10)
    link(doc, "GitHub 项目与使用入口", "https://github.com/zez2001/LabSim/tree/codex/labsim")
    link(doc, "teacher_full.mp4  完整teacher装载", "teacher_full.mp4")
    link(doc, cfg.get("policy_video_label", "gpt6_astra_incomplete.mp4  GPT真实接入 任务未完成"), videos[1][1])
    heading(doc, "剩余问题与下一步")
    for text in cfg.get("remaining", [
        "完成同一Isaac进程内成功—受控结束—reset后成功，并核查旧ID失效、接触和控制状态清空、多集索引及文件句柄闭合。",
        "通过真实已部署外部模型完成持续RPC交互及正常终止，分别报告任务结果与接口、服务或环境异常。",
        "测量完整teacher的启动、rollout、写入关闭、读取及RAM/VRAM开销；只实施有同口径证据的记录与存储优化。",
    ]):
        para(doc, text, size=10)
    para(doc, cfg.get("release_note", "本页待验收项以2026年10月27日已有闭合证据为基线；本轮环境验收完成后更新。汇报只说明仿真环境能力，未开展训练、随机场景采集或真机验证。"), size=9)

    dest = out / NAME
    doc.save(dest)
    receipt = {"schema": "labsim.Project_release_v1.meeting_build.v1", "docx": str(dest),
        "docx_sha256": sha(dest), "media": media, "expected_pages": 4,
        "package_allowed_files": [NAME] + [x[1] for x in videos],
        "fonts_declared": {"Chinese_body": "宋体", "Chinese_heading": "黑体", "Latin_digits": "Arial"},
        "render_font_availability": "Host lacks requested proprietary fonts; existing CJK/Liberation fallbacks used for QA only.",
        "marker_already_run": True, "overrides": str(a.overrides) if a.overrides else None}
    (HERE / "meeting_build_receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2)+"\n")
    print(json.dumps({"docx": str(dest), "files": len(list(out.iterdir())),
                      "bytes": sum(x.stat().st_size for x in out.iterdir())}, ensure_ascii=False))


if __name__ == "__main__":
    main()
